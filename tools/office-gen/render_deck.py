#!/usr/bin/env python3
"""
Deck spec (JSON) -> .pptx. Every pixel decided here, not by the model.

THE BET. A 4B model is good at prose and at picking from a short menu. It is bad
at spatial reasoning, EMU arithmetic, colour harmony and holding a grid across a
dozen slides. Asking it for python-pptx puts all four weaknesses on the critical
path. So it emits CONTENT plus a layout name; this file owns the design. The deck
looks the same whether the spec came from a 4B or a frontier model — only the
words change.

V2, after looking at V1 and agreeing it read "template basic":
  - asymmetry everywhere (colour fields, off-centre content) instead of centred
    boxes with a rule under the title;
  - scale contrast — 96pt numerals against 12pt captions;
  - data visualisation drawn as native shapes (viz.py) rather than one default
    bar chart, because a default chart looks like a default chart;
  - a motif (the accent bar) repeated so the deck reads as a set.

VQ-01 (2026-09-23), the correctness pass — what the renderer did to ordinary
input, measured by tools/visual-eval on replays of real 4B output:
  - numbers parse as written ("22M", "$38k", "3,100" — numparse.py), and a
    display that is only a unit ("GW") keeps its number;
  - every block of text that something is stacked under is MEASURED (textfit,
    with the rendered line height), including numerals, which shrink to their
    column instead of breaking "3,10 / 0" over their label;
  - "emphasis": "false" is false;
  - nothing is dropped silently: a list longer than a layout holds is cut AND
    reported ("5 of 10 items shown"), a layout that fails leaves no half-built
    slide behind and says so, a blank table header no longer crashes;
  - `build` reports what it DREW (`drawn=`), which is what office.py tells the
    model — not what the spec asked for.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).parent))
import viz  # noqa: E402
import hero  # noqa: E402
import numparse  # noqa: E402
import textfit  # noqa: E402

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.9)
GUT = Inches(0.5)
CW = W - 2 * M
FONT = "Helvetica Neue"
# A serif for the one quiet slide. The benchmark deck set its pull quote in
# serif italic and it read as editorial rather than as a slide; the same
# sentence in bold sans reads as a heading. Georgia ships on every Mac.
SERIF = "Georgia"
EMU_PT = 12700


class Theme:
    def __init__(self, primary, support, accent, deep=None,
                 ink="#141419", paper="#FCFBF8", mute="#75747E", faint="#E7E5DF"):
        c = lambda s: RGBColor.from_string(s.lstrip("#"))
        self.primary, self.support, self.accent = c(primary), c(support), c(accent)
        self.deep = c(deep or primary)
        self.ink, self.paper, self.mute, self.faint = c(ink), c(paper), c(mute), c(faint)


THEMES = {
    "midnight": Theme("1E2761", "CADCFC", "F26B5B", deep="141A3F"),
    "forest": Theme("22452B", "9CBF7E", "E4A045", deep="16301C"),
    "terracotta": Theme("8C3D2B", "E7E1D3", "2F6B6B", deep="63291D"),
    "charcoal": Theme("24262C", "D9DCE1", "E2603F", deep="15171B"),
    "ink": Theme("14203C", "C7D3E8", "D98E32", deep="0C1428"),
}


# ── what was drawn ───────────────────────────────────────────────────────────
class SlideReport:
    """What one slide ended up showing — the truth office.py reports back."""

    def __init__(self, index: int, layout: str):
        self.index = index
        self.layout = layout
        self.asked = layout
        self.head = ""
        self.tail = ""
        self.items: list[str] = []
        self.cuts: list[tuple[str, int, int]] = []     # (field, given, shown)
        self.notes: list[str] = []

    def as_dict(self) -> dict:
        return {"slide": self.index, "layout": self.layout, "asked": self.asked,
                "head": self.head, "tail": self.tail, "items": self.items,
                "cuts": [{"field": f, "given": g, "shown": s} for f, g, s in self.cuts],
                "notes": self.notes}


_R: SlideReport | None = None


def _cap(seq, n: int, field: str) -> list:
    """The first n items — and a record of it when that is not all of them."""
    items = _list(seq)
    if len(items) > n and _R is not None:
        _R.cuts.append((field, len(items), n))
    return items[:n]


def _note(msg: str) -> None:
    if _R is not None:
        _R.notes.append(msg)


# ── coercion ─────────────────────────────────────────────────────────────────
def _s(v) -> str:
    """Whatever the model produced -> a string. A 4B hands you a list where the
    schema said string often enough that this is load-bearing, not defensive."""
    return numparse.as_text(v)


def _list(v) -> list:
    if v is None:
        return []
    if isinstance(v, (list, tuple)):
        return list(v)
    if isinstance(v, dict):
        return list(v.values())
    return [v]


def _num(v, default=0.0) -> float:
    return numparse.num(v, default)


def _get(v, key, default=None):
    """A field of an item the model may have written as an object OR a bare value."""
    if isinstance(v, dict):
        return v.get(key, default)
    return default


def _heading(v) -> str:
    """A column's or side's heading, whatever the model called it."""
    if isinstance(v, dict):
        return _s(v.get("heading") or v.get("title") or v.get("name") or v.get("label") or "")
    return _s(v)


# ── primitives ───────────────────────────────────────────────────────────────
def _rect(sl, x, y, w, h, colour, *, radius=None):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    s = sl.shapes.add_shape(shape, int(x), int(y), int(max(w, 1)), int(max(h, 1)))
    s.fill.solid()
    s.fill.fore_color.rgb = colour
    s.line.fill.background()
    s.shadow.inherit = False
    if radius is not None:
        s.adjustments[0] = radius
    return s


def _tf(sl, x, y, w, h, *, wrap=True):
    tb = sl.shapes.add_textbox(int(x), int(y), int(w), int(h))
    f = tb.text_frame
    f.word_wrap = wrap
    f.margin_left = f.margin_right = f.margin_top = f.margin_bottom = 0
    return f


def _p(tf, text, size, colour, *, bold=False, align=PP_ALIGN.LEFT, after=0,
       line=None, first=False, spacing=None, font=None, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(after)
    if line:
        p.line_spacing = line
    r = p.add_run()
    r.text = _s(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = font or FONT
    r.font.color.rgb = colour
    if italic:
        r.font.italic = True
    if spacing is not None:
        # Letter-spacing has no python-pptx API; set it on the run properties.
        r.font._rPr.set("spc", str(int(spacing * 100)))
    return p


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _bg(sl, colour):
    return _rect(sl, 0, 0, W, H, colour)


def _kicker(sl, t, text, x=None, y=None, colour=None):
    """The small label over a title. Measured with its letter-spacing, and held
    to two lines — a 4B writes a sentence here often enough (VQ-08 lints it)."""
    if not text:
        return y or Inches(0.72)
    top = y if y is not None else Inches(0.72)
    end, _sz, _n = _block(sl, _s(text).upper(), x or M, top, CW / Inches(1), [10.5], 2,
                          colour or t.accent, bold=True, spacing=1.4, line=1.1)
    return end


def _fit(text, box_inches, sizes, max_lines, *, bold=False, font=FONT, italic=False, spacing=None):
    """Largest size that fits, plus the exact lines. Real metrics, not a guess."""
    return textfit.fit_size(_s(text), box_inches * 72, sizes, max_lines,
                            font=font, bold=bold, italic=italic, tracking=spacing or 0.0)


def _block(sl, text, x, y, box_inches, sizes, max_lines, colour, *, bold=False,
           font=FONT, line=1.05, spacing=None, italic=False, align=PP_ALIGN.LEFT):
    """Measure, wrap, then write each line as its own paragraph with reflow OFF.
    What we measured is exactly what renders. Returns (bottom, size, lines)."""
    size, lines = _fit(text, box_inches, sizes, max_lines, bold=bold, font=font,
                       italic=italic, spacing=spacing)
    h = Inches(textfit.height_pt(len(lines), size, line) / 72 + 0.06)
    tf = _tf(sl, x, y, Inches(box_inches), h, wrap=False)
    for i, ln in enumerate(lines):
        _p(tf, ln, size, colour, bold=bold, first=(i == 0), line=line,
           font=font, spacing=spacing, italic=italic, align=align)
    return y + h, size, len(lines)


def _drawn(lines_text: str) -> str:
    return " ".join(lines_text.split())


def _head(sl, t, s, *, y=Inches(1.28), width=0.82):
    kick_end = _kicker(sl, t, s.get("kicker"))
    y = max(y, kick_end + Inches(0.14))
    title = _s(s.get("title"))
    end, size, _n = _block(sl, title, M, y, 11.5 * width,
                           [34, 30, 26, 22], 2, t.ink, bold=True, line=1.04)
    if _R is not None:
        _R.head = _drawn(" ".join(textfit.wrap(title, 11.5 * width * 72, size, bold=True, max_lines=2)))
    ry = end + Inches(0.10)
    _rect(sl, M, ry, Inches(0.62), Pt(4), t.accent)
    return ry + Inches(0.5)


def _stack(sl, texts, x, y, box_inches, bottom, sizes, colour, *, line=1.35, gap=Inches(0.28),
           bullet=None, bold=False, font=FONT):
    """Paragraph-per-item text stacked by MEASURED height, the size stepped down
    until the whole list fits above `bottom`. Returns the y after the last item."""
    items = [_s(b) for b in texts if _s(b).strip()]
    if not items:
        return y
    for size in sizes:
        need = sum(Inches(textfit.height_pt(len(textfit.wrap(b, box_inches * 72, size, bold=bold, font=font)),
                                            size, line) / 72) for b in items) + gap * (len(items) - 1)
        if y + need <= bottom:
            break
    cy = y
    for b in items:
        end, _sz, _n = _block(sl, b, x, cy, box_inches, [size], 4, colour, line=line, bold=bold, font=font)
        if bullet is not None:
            bullet(cy, size)
        cy = end + gap
    return cy - gap


# ── layouts ──────────────────────────────────────────────────────────────────
def L_title(prs, t, s):
    sl = _blank(prs)
    _bg(sl, t.deep)
    # Asymmetric colour field off the right edge. This one shape is most of the
    # difference between "designed" and "default".
    _rect(sl, W * 0.62, 0, W * 0.38, H, t.primary)
    _rect(sl, W * 0.62, 0, Inches(0.09), H, t.accent)
    _rect(sl, M, Inches(2.28), Inches(1.3), Pt(7), t.accent)
    title = _s(s.get("title"))
    end, _sz, _n = _block(sl, title, M, Inches(2.72), 7.3,
                          [52, 46, 40, 34], 3, t.paper, bold=True, line=1.04)
    if _R is not None:
        _R.head = title
    foot_top = H - Inches(0.95)
    if s.get("subtitle"):
        sub_y = end + Inches(0.3)
        room = max(1, int((foot_top - Inches(0.15) - sub_y) / Pt(textfit.line_pitch_pt(15, 1.3))))
        _block(sl, s["subtitle"], M, sub_y, (W * 0.48) / Inches(1), [17, 15.5, 14], min(3, room),
               t.support, line=1.3)
        if _R is not None:
            _R.tail = _s(s["subtitle"])
    if s.get("footnote"):
        tf3 = _tf(sl, M, foot_top, W * 0.5, Inches(0.3))
        _p(tf3, s["footnote"], 10.5, t.support, first=True, spacing=0.8)
    return sl


def L_section(prs, t, s):
    sl = _blank(prs)
    _bg(sl, t.deep)
    # Oversized numeral — scale contrast as the whole idea of the slide.
    num = _s(s.get("number") or "")
    if num:
        tf0 = _tf(sl, M - Inches(0.22), Inches(1.4), Inches(5.4), Inches(3.4))
        _p(tf0, str(num).zfill(2), 150, t.primary, bold=True, line=0.9, first=True)
    _rect(sl, M, Inches(4.62), Inches(1.3), Pt(7), t.accent)
    t_top = Inches(4.95)
    t_end, _sz, _n = _block(sl, s.get("title"), M, t_top, 8.6,
                            [36, 31, 27], 2, t.paper, bold=True, line=1.05)
    if _R is not None:
        _R.head = _s(s.get("title"))
    if s.get("summary"):
        # BELOW the title, then SHRUNK until it fits above the bottom margin.
        # The previous attempt clamped its y upward instead, which slid a long
        # summary back into the title — a worse bug than the overflow it fixed.
        s_top = t_end + Inches(0.18)
        pitch = Pt(textfit.line_pitch_pt(11.5, 1.32))
        max_lines = max(1, int((H - Inches(0.45) - s_top) / pitch))
        _block(sl, s["summary"], M, s_top, 9.4, [13.5, 12.5, 11.5, 10.5],
               min(4, max_lines), t.support, line=1.32)
        if _R is not None:
            _R.tail = _s(s["summary"])
    return sl


def L_bullets(prs, t, s):
    sl = _blank(prs)
    _bg(sl, t.paper)
    _rect(sl, 0, 0, Inches(0.16), H, t.primary)   # spine motif
    y = _head(sl, t, s)
    bullets = _cap(s.get("bullets"), 5, "bullets")
    if _R is not None:
        _R.items = [_s(b) for b in bullets]

    def dot(cy, size):
        _rect(sl, M, cy + Pt(size * 0.585), Inches(0.2), Inches(0.2), t.accent, radius=0.5)

    _stack(sl, bullets, M + Inches(0.5), y, (CW * 0.78) / Inches(1), H - Inches(0.75),
           [16, 15, 14], t.ink, line=1.35, gap=Inches(0.26), bullet=dot)
    return sl


def L_two_column(prs, t, s):
    sl = _blank(prs)
    _bg(sl, t.paper)
    y = _head(sl, t, s)
    cw = (CW - GUT) / 2
    cols = _cap(s.get("columns"), 2, "columns")
    for i, col in enumerate(cols):
        x = M + i * (cw + GUT)
        _rect(sl, x, y, Inches(0.38), Pt(4), t.accent if i == 0 else t.support)
        end, _sz, _n = _block(sl, _heading(col), x, y + Inches(0.2), cw / Inches(1),
                              [17, 15.5], 2, t.primary, bold=True, line=1.1)
        points = _cap(_get(col, "points", []), 5, f"points in column {i + 1}")
        _stack(sl, points, x, end + Inches(0.18), cw / Inches(1), H - Inches(0.75),
               [14, 13, 12], t.ink, line=1.45, gap=Inches(0.18))
    return sl


def L_stats(prs, t, s):
    """Oversized numerals on a colour field. No cards — the numbers are the
    design."""
    sl = _blank(prs)
    _bg(sl, t.paper)
    _rect(sl, 0, H * 0.46, W, H * 0.54, t.deep)
    kick_end = _kicker(sl, t, s.get("kicker"))
    title = _s(s.get("title"))
    _block(sl, title, M, max(Inches(1.24), kick_end + Inches(0.14)), 9.2, [34, 30, 27, 24], 2,
           t.ink, bold=True, line=1.05)
    if _R is not None:
        _R.head = title
    items = _cap(s.get("stats"), 4, "stats")
    n = max(1, len(items))
    cw = CW / n
    values = [_s(_get(it, "value", it)) for it in items]
    labels = [_s(_get(it, "label", "")) for it in items]
    # One numeral size for the row — the largest at which EVERY value sits on
    # one line in its column. A fixed 76 pt broke "3,100" into "3,10 / 0" over
    # its own label (D2).
    box = (cw - Inches(0.3)) / Inches(1) * 72
    size = 76
    for cand in (76, 68, 60, 54, 48, 42, 36, 32, 28):
        size = cand
        if all(textfit.width_pt(v, cand, bold=True) <= box for v in values):
            break
    rule_y = H * 0.46 + Inches(1.92)
    num_h = Pt(textfit.height_pt(1, size, 0.95))
    for i, (value, label) in enumerate(zip(values, labels)):
        x = M + i * cw
        # Sit the numeral on the rule, whatever size it came out at.
        tf1 = _tf(sl, x, rule_y - Inches(0.12) - num_h, cw - Inches(0.3), num_h, wrap=False)
        _p(tf1, value, size, t.paper, bold=True, line=0.95, first=True)
        _rect(sl, x, rule_y, Inches(0.5), Pt(4), t.accent)
        _block(sl, label, x, H * 0.46 + Inches(2.14), (cw - Inches(0.45)) / Inches(1),
               [12.5, 11.5], 4, t.support, line=1.35)
    if _R is not None:
        _R.items = [f"{v} {lab}".strip() for v, lab in zip(values, labels)]
    if s.get("note"):
        tf3 = _tf(sl, M, H * 0.46 - Inches(0.62), CW * 0.75, Inches(0.5))
        _p(tf3, s["note"], 12.5, t.mute, line=1.3, first=True)
    return sl


def L_comparison(prs, t, s):
    """Split screen, full bleed. The contrast IS the comparison, so no cards."""
    sl = _blank(prs)
    _bg(sl, t.paper)
    y = _head(sl, t, s)
    half = W / 2
    _rect(sl, 0, y - Inches(0.28), half, H - y + Inches(0.28), t.deep)
    _rect(sl, half, y - Inches(0.28), half, H - y + Inches(0.28), t.support)
    for i, side in enumerate(_cap(s.get("sides"), 2, "sides")):
        x = (M if i == 0 else half + Inches(0.7))
        fg = t.paper if i == 0 else t.ink
        end, _sz, _n = _block(sl, _heading(side), x, y + Inches(0.18),
                              (half - Inches(1.5)) / Inches(1), [19, 17], 2, fg, bold=True, line=1.1)
        _rect(sl, x, end + Inches(0.12), Inches(0.42), Pt(4),
              t.accent if i == 0 else t.primary)
        points = _cap(_get(side, "points", []), 5, f"points on side {i + 1}")
        _stack(sl, points, x, end + Inches(0.4), (half - Inches(1.6)) / Inches(1), H - Inches(0.5),
               [13.5, 12.5, 11.5], fg, line=1.45, gap=Inches(0.16))
    return sl


def L_bars(prs, t, s):
    """Horizontal bars drawn as shapes — the data-viz workhorse."""
    sl = _blank(prs)
    _bg(sl, t.paper)
    y = _head(sl, t, s)
    given = _list(s.get("items"))
    raw = _cap(given, 5, "items")
    items = []
    for it in raw:
        value = _get(it, "value", it)
        parsed = numparse.parse(value)
        if parsed is None:
            _note(f"'{_s(value)}' is not a number, drawn as an empty bar")
        items.append({"label": _s(_get(it, "label", "")),
                      "value": parsed.value if parsed is not None else 0.0,
                      "display": numparse.shown(value, _get(it, "display"))})
    raw_hi = s.get("highlight_index")
    if isinstance(raw_hi, bool) or raw_hi in (None, ""):
        hi = 0                     # unspecified: the first bar leads, as it always did
    elif isinstance(raw_hi, (int, float)) or (isinstance(raw_hi, str) and raw_hi.strip().isdigit()):
        hi = int(raw_hi)
        if not 0 <= hi < len(items):
            _note(f"highlight_index {hi} is not among the {len(items)} bars shown; nothing highlighted")
            hi = None
    else:
        hi = None
    if items:
        note_room = Inches(0.9) if s.get("note") else Inches(0.35)
        viz.bar_rows(sl, t, M, y + Inches(0.15), CW * 0.88, items,
                     avail=H - M - y - note_room, highlight=hi)
    if _R is not None:
        _R.items = [f"{i['label']} {i['display']}".strip() for i in items]
    if s.get("note"):
        tf = _tf(sl, M, H - M - Inches(0.55), CW * 0.8, Inches(0.5))
        _p(tf, s["note"], 12, t.mute, line=1.3, first=True)
    return sl


def L_waffle(prs, t, s):
    """Proportion as a 10x10 dot grid, with the claim set large beside it."""
    sl = _blank(prs)
    _bg(sl, t.paper)
    y = _head(sl, t, s, width=0.62)
    pct = max(0, min(100, int(round(_num(s.get("percent"), 0)))))
    viz.waffle(sl, t, W - M - Inches(3.4), Inches(1.9), pct)
    tf = _tf(sl, M, y + Inches(0.25), CW * 0.5, Inches(1.6))
    _p(tf, f"{pct}%", 96, t.primary, bold=True, line=0.95, first=True)
    _rect(sl, M, y + Inches(1.62), Inches(0.62), Pt(4), t.accent)
    _block(sl, s.get("label"), M, y + Inches(1.88), (CW * 0.46) / Inches(1), [16, 14.5, 13], 4,
           t.ink, line=1.4)
    if _R is not None:
        _R.items = [f"{pct}% {_s(s.get('label'))}".strip()]
    if s.get("note"):
        tf3 = _tf(sl, M, H - M - Inches(0.5), CW * 0.55, Inches(0.5))
        _p(tf3, s["note"], 12, t.mute, line=1.3, first=True)
    return sl


def L_range(prs, t, s):
    """Labelled bands on one axis — for threshold claims, where two numbers in a
    sentence hide the shape of the finding."""
    sl = _blank(prs)
    _bg(sl, t.paper)
    y = _head(sl, t, s)
    bands = [{"label": _s(_get(b, "label", "")), "from": _num(_get(b, "from")),
              "to": _num(_get(b, "to"), 1)} for b in _cap(s.get("bands"), 4, "bands")]
    end = y
    if bands:
        note_room = Inches(1.1) if s.get("note") else Inches(0.4)
        end = viz.range_bands(sl, t, M, y + Inches(0.62), CW * 0.86, bands,
                              avail=H - M - y - Inches(0.62) - note_room)
    if _R is not None:
        _R.items = [b["label"] for b in bands]
    if s.get("note"):
        tf = _tf(sl, M, min(end + Inches(0.2), H - M - Inches(0.9)), CW * 0.82, Inches(0.9))
        _p(tf, s["note"], 13.5, t.ink, line=1.4, first=True)
    return sl


def _cell_text(cell, text, size, colour, *, bold=False):
    """Set a cell's text AND style it — an empty cell has no run to style, which
    is how a blank corner header crashed the whole table (D8)."""
    cell.text = _s(text)
    p = cell.text_frame.paragraphs[0]
    r = p.runs[0] if p.runs else p.add_run()
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
    r.font.color.rgb = colour


def L_table(prs, t, s):
    sl = _blank(prs)
    _bg(sl, t.paper)
    y = _head(sl, t, s)
    rows = [_list(r) for r in _cap(s.get("rows"), 6, "rows")]
    cols = [_s(c) for c in _list(s.get("headers"))]
    width = max([len(cols)] + [len(r) for r in rows] + [1])
    if len(cols) < width:
        # A row wider than the header: keep every cell, the header just has
        # nothing to say over the extra column.
        cols += [""] * (width - len(cols))
    avail = H - M - y - Inches(0.5)
    row_h = max(Inches(0.5), min(Inches(0.95), avail / (len(rows) + 1)))
    shape = sl.shapes.add_table(len(rows) + 1, len(cols), int(M), int(y),
                                int(CW), int(row_h * (len(rows) + 1)))
    tbl = shape.table
    tbl.first_row = True
    for c, name in enumerate(cols):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = t.deep
        _cell_text(cell, name, 12, t.paper, bold=True)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for r, row in enumerate(rows, start=1):
        for c in range(len(cols)):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = t.paper if r % 2 else t.faint
            _cell_text(cell, row[c] if c < len(row) else "", 12, t.ink, bold=(c == 0))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if _R is not None:
        _R.items = [_s(r[0]) if r else "" for r in rows]
    return sl


def L_quote(prs, t, s):
    sl = _blank(prs)
    _bg(sl, t.deep)
    _rect(sl, 0, 0, Inches(0.16), H, t.accent)
    tf0 = _tf(sl, M, Inches(1.3), Inches(2.0), Inches(1.8))
    _p(tf0, "“", 150, t.primary, bold=True, first=True)
    q = _s(s.get("quote"))
    # Measured in the serif italic it is set in, held above the attribution.
    top, floor = Inches(2.5), Inches(5.3)
    size = 25
    for cand in (38, 34, 31, 28, 25, 22):
        size = cand
        n = len(textfit.wrap(q, (CW * 0.84) / Inches(1) * 72, cand, font=SERIF, italic=True))
        if top + Pt(textfit.height_pt(n, cand, 1.26)) <= floor:
            break
    lines = int((floor - top) / Pt(textfit.line_pitch_pt(size, 1.26)))
    _block(sl, q, M, top, (CW * 0.84) / Inches(1), [size], max(1, lines), t.paper,
           font=SERIF, italic=True, line=1.26)
    if _R is not None:
        _R.head = q
    if s.get("attribution"):
        _rect(sl, M, Inches(5.55), Inches(0.5), Pt(4), t.accent)
        tf2 = _tf(sl, M, Inches(5.8), CW * 0.6, Inches(0.5))
        _p(tf2, _s(s["attribution"]), 13, t.support, first=True)
        if _R is not None:
            _R.tail = _s(s["attribution"])
    return sl


def L_closing(prs, t, s):
    sl = _blank(prs)
    _bg(sl, t.deep)
    _rect(sl, W * 0.60, 0, W * 0.40, H, t.primary)
    title = _s(s.get("title"))
    end, _sz, _n = _block(sl, title, M, Inches(1.55), (W * 0.5) / Inches(1), [40, 36, 32, 28], 2,
                          t.paper, bold=True, line=1.05)
    if _R is not None:
        _R.head = title
    rule_y = max(Inches(3.0), end + Inches(0.2))
    _rect(sl, M, rule_y, Inches(1.3), Pt(7), t.accent)
    points = _cap(s.get("points"), 4, "points")
    cy = rule_y + Inches(0.42)
    size = 15
    for cand in (15, 14, 13):
        size = cand
        need = sum(Inches(textfit.height_pt(len(textfit.wrap(_s(b), (W * 0.46) / Inches(1) * 72, cand)),
                                            cand, 1.35) / 72) + Inches(0.3) for b in points)
        if cy + need <= H - Inches(0.6):
            break
    for i, b in enumerate(points):
        tf2 = _tf(sl, M, cy, Inches(0.42), Inches(0.5))
        _p(tf2, f"0{i+1}", size, t.accent, bold=True, first=True)
        end_b, _sz, _n = _block(sl, b, M + Inches(0.62), cy, (W * 0.46) / Inches(1), [size], 4,
                                t.support, line=1.35)
        cy = end_b + Inches(0.3)
    if _R is not None:
        _R.items = [_s(b) for b in points]
    return sl


def L_flow(prs, t, s):
    """A process, drawn. The benchmark's render-look-fix diagram was the single
    clearest slide in it."""
    sl = _blank(prs)
    _bg(sl, t.paper)
    y = _head(sl, t, s)
    raw = _cap(s.get("steps"), 5, "steps")
    # Emphasis only where the model said TRUE ("false" was truthy — D5). When no
    # step says anything either way, the first step is the start marker, as the
    # layout always drew it.
    said = any(isinstance(x, dict) and "emphasis" in x for x in raw)
    steps = [{"label": _s(_get(x, "label", x)), "caption": _s(_get(x, "caption", "")),
              "emphasis": numparse.truthy(_get(x, "emphasis")) if said else i == 0}
             for i, x in enumerate(raw)]
    if steps:
        top = y + Inches(0.5)
        if not s.get("note"):
            # No note band under it: the diagram is centred in the body instead
            # of hugging the title over an empty lower half. MEASURED on the
            # Tidewell replay: once the invented note line was stripped (VQ-08)
            # the band went with it and 37.5% of the slide stood empty.
            block = Inches(0.62) * 2 + (Inches(1.02) if any(st["caption"] for st in steps) else 0)
            room = H - Inches(0.9) - top
            if room > block:
                top += (room - block) / 2
        viz.flow(sl, t, M, top, CW, steps)
    if _R is not None:
        _R.items = [st["label"] for st in steps]
    if s.get("note"):
        _rect(sl, 0, H - Inches(1.15), W, Inches(1.15), t.deep)
        tf = _tf(sl, M, H - Inches(0.86), CW * 0.85, Inches(0.6))
        _p(tf, s["note"], 17, t.paper, bold=True, line=1.25, first=True)
    return sl


def L_matrix(prs, t, s):
    """Claims against options, with filled cells. Replaces two bullet lists."""
    sl = _blank(prs)
    _bg(sl, t.paper)
    y = _head(sl, t, s)
    cols = [_s(c) for c in _cap(s.get("columns"), 3, "columns")] or ["A", "B"]
    rows = []
    for r in _cap(s.get("rows"), 6, "rows"):
        # The model returns rows as objects OR as bare arrays like
        # ["claim", true, false]. Both are reasonable readings of the schema, so
        # accept both rather than dropping the slide.
        if isinstance(r, dict):
            vals = _list(r.get("values"))
            rows.append({"label": _s(r.get("label")),
                         "values": [numparse.truthy(v, loose=True) for v in vals]})
        else:
            arr = _list(r)
            rows.append({"label": _s(arr[0] if arr else ""),
                         "values": [numparse.truthy(v, loose=True) for v in arr[1:]]})
    if rows:
        viz.matrix(sl, t, M, y + Inches(0.55), CW, rows, cols)
    if _R is not None:
        _R.items = [r["label"] for r in rows]
    if s.get("verdict"):
        _rect(sl, 0, H - Inches(1.05), W, Inches(1.05), t.deep)
        tf = _tf(sl, M, H - Inches(0.78), CW * 0.7, Inches(0.55))
        _p(tf, s["verdict"], 15, t.paper, line=1.3, first=True)
    return sl


def _hero(fn):
    def draw(prs, t, s):
        if fn is hero.hero_split and s.get("body") is not None:
            s = {**s, "body": _cap(s.get("body"), 4, "body paragraphs")}
        sl = fn(prs, t, s)
        if _R is not None:
            _R.head = _s(s.get("title") or s.get("statement"))
            _R.tail = _s(s.get("subtitle") or s.get("attribution") or "")
            if s.get("body"):
                _R.items = [_s(b) for b in _list(s.get("body"))[:4]]
        return sl
    return draw


LAYOUTS = {
    "flow": L_flow,
    "matrix": L_matrix,
    # High-fidelity grounds. Registered ahead of the flat layouts so a spec can
    # opt into imagery per slide.
    "hero_title": _hero(hero.hero_title),
    "hero_split": _hero(hero.hero_split),
    "hero_statement": _hero(hero.hero_statement),
    "title": L_title, "section": L_section, "bullets": L_bullets,
    "two_column": L_two_column, "stats": L_stats, "comparison": L_comparison,
    "bars": L_bars, "waffle": L_waffle, "range": L_range,
    "table": L_table, "quote": L_quote, "closing": L_closing,
}
DARK = {"title", "section", "quote", "closing", "stats",
        "hero_title", "hero_statement", "flow", "matrix"}


def _footer(sl, t, text, page):
    tf = _tf(sl, M, H - Inches(0.52), CW * 0.7, Inches(0.28))
    _p(tf, text, 9, t.mute, first=True, spacing=0.6)
    tf2 = _tf(sl, W - M - Inches(0.6), H - Inches(0.52), Inches(0.6), Inches(0.28))
    _p(tf2, str(page), 9, t.mute, align=PP_ALIGN.RIGHT, first=True)


def _drop_slides_from(prs, n: int) -> None:
    """Remove every slide from index n on — the half-built slide a failed layout
    left behind (D7). Dropping the relationship drops the part from the save."""
    lst = prs.slides._sldIdLst
    for sid in list(lst)[n:]:
        prs.part.drop_rel(sid.rId)
        lst.remove(sid)


def _as_bullets(s: dict) -> list[str]:
    """Everything a slide said, as lines — so a layout that cannot be drawn
    still shows its content rather than only its title."""
    out: list[str] = []
    for key in ("bullets", "points", "body"):
        out += [_s(x) for x in _list(s.get(key)) if _s(x).strip()]
    for key, fields in (("items", ("label", "display", "value")), ("steps", ("label", "caption")),
                        ("stats", ("value", "label")), ("bands", ("label",)), ("rows", None)):
        for it in _list(s.get(key)):
            if fields is None:
                out.append("  ·  ".join(_s(c) for c in _list(it) if _s(c).strip()))
            elif isinstance(it, dict):
                out.append(" — ".join(_s(it.get(f)) for f in fields if _s(it.get(f)).strip()))
            else:
                out.append(_s(it))
    for side in _list(s.get("sides")) + _list(s.get("columns")):
        if isinstance(side, dict):
            out += [_s(p) for p in _list(side.get("points"))]
    for key in ("statement", "quote", "summary", "note", "verdict", "subtitle"):
        if _s(s.get(key)).strip():
            out.append(_s(s[key]))
    return [x for x in out if x.strip()] or [_s(s.get("title"))]


def build(spec: dict, out: Path, drawn: list | None = None, warnings: list | None = None) -> Path:
    """Render the deck. `drawn` (if given) receives one record per slide of what
    was actually drawn; `warnings` receives one line per thing the model should
    know (an item cut, a layout that failed, a highlight that points nowhere)."""
    global _R
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    t = THEMES.get(spec.get("theme", "ink"), THEMES["ink"])
    running = _s(spec.get("running_title"))
    for i, s in enumerate(spec.get("slides", [])):
        if not isinstance(s, dict):
            s = {"layout": "bullets", "title": _s(s)}
        layout = _s(s.get("layout")) or "bullets"
        _R = SlideReport(i + 1, layout)
        fn = LAYOUTS.get(layout)
        if fn is None:
            _R.notes.append(f"layout '{layout}' does not exist; drawn as bullets")
            s = {**s, "bullets": _as_bullets(s)}
            fn, layout = L_bullets, "bullets"
        before = len(prs.slides)
        try:
            sl = fn(prs, t, s)
        except Exception as err:  # noqa: BLE001
            # A malformed slide degrades to bullets rather than losing the deck —
            # and takes its half-built slide with it.
            _drop_slides_from(prs, before)
            print(f"  ! slide {i+1} ({layout}) failed: {err} — drew it as bullets", file=sys.stderr)
            asked = layout
            _R = SlideReport(i + 1, "bullets")
            _R.asked = asked
            _R.notes.append(f"could not be drawn as {asked} ({type(err).__name__}: {err}); drawn as bullets")
            sl = L_bullets(prs, t, {**s, "bullets": _as_bullets(s)})
            layout = "bullets"
        _R.layout = layout
        if i > 0 and layout not in DARK:
            _footer(sl, t, running, i + 1)
        if drawn is not None:
            drawn.append(_R.as_dict())
        if warnings is not None:
            for field, given, shown in _R.cuts:
                warnings.append(f"slide {i+1} ({layout}): {field} — {shown} of {given} shown; "
                                f"the layout holds {shown}, split the rest onto another slide")
            for note in _R.notes:
                warnings.append(f"slide {i+1} ({_R.asked}): {note}")
    _R = None
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


if __name__ == "__main__":
    spec_path = Path(sys.argv[1])
    out_path = Path(sys.argv[2]) if len(sys.argv) > 2 else spec_path.with_suffix(".pptx")
    spec = json.loads(spec_path.read_text())
    warn: list[str] = []
    p = build(spec, out_path, warnings=warn)
    print(f"wrote {p}  ({len(spec.get('slides', []))} slides)")
    for w in warn:
        print(f"  ! {w}")
