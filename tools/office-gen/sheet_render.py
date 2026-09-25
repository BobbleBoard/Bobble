#!/usr/bin/env python3
"""
Spec (JSON) -> .xlsx that looks like a dashboard, not a data dump.

A generated spreadsheet is usually raw values in A1 with default column widths,
which is legible to a machine and hostile to a person. Everything that makes a
workbook look considered — column widths, a title band, KPI cells, banded rows,
number formats, frozen panes, a native chart — is mechanical and none of it is
something a model should be spending tokens on.

Formulas are preserved as formulas. Writing a computed VALUE where a SUM belongs
produces a workbook that looks right and dies the moment anyone edits a cell.

WF-06 (2026-09-24), the render path's workbook (`office.py render xlsx`, a spec
built by code from pages it fetched — never by the make path's model). Given
`cite=` it may have:
  - several sheets: `{"sheets": [<a sheet spec, as below>, …]}` (a research
    workbook is Findings, Data and Sources, workflows.md §4.5);
  - cells that are more than a value — {"text"|"value": …} plus any of
    "url" (a live link), "note" (a cell note a reader hovers) and
    "cite": ["S3", "S7"] (shown "3, 7", linked to source 3's row);
  - a Sources sheet, from the spec's sources: every URL a live link.
Links are the cell's own hyperlink, not a HYPERLINK() formula: openpyxl writes
a formula with no cached value, so QuickLook, Numbers and any viewer that does
not recalculate would show an empty cell where the link should be. Without
`cite` (every make) none of this runs, and a single sheet is drawn exactly as
it always was — a small model's spec does not get to add links (VQ-08).
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, DataBarRule  # noqa: F401
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

sys.path.insert(0, str(Path(__file__).parent))
import palette as pal  # noqa: E402

FONT = "Helvetica Neue"


def _hx(h: str) -> str:
    return h.lstrip("#").upper()


def _fill(h: str) -> PatternFill:
    return PatternFill("solid", fgColor=_hx(h))


def _font(t, *, size=11, bold=False, colour=None, italic=False) -> Font:
    return Font(name=FONT, size=size, bold=bold, italic=italic,
                color=_hx(colour or t.ink))


def _band(ws, row, col0, col1, t, colour, height=None):
    for c in range(col0, col1 + 1):
        ws.cell(row=row, column=c).fill = _fill(colour)
    if height:
        ws.row_dimensions[row].height = height


class _Plain:
    """How the make path writes a cell: the value, as given (a header as its
    string) — exactly what this file always wrote."""

    @staticmethod
    def put(ws, row, col, v):
        return ws.cell(row=row, column=col, value=v)

    @staticmethod
    def head(h):
        return str(h)

    @staticmethod
    def shown(v) -> str:
        return str(v)

    @staticmethod
    def text(v):
        return v

    def style(self, cell, t):
        pass


def _sheet(ws, spec: dict, t, cells=_Plain()) -> dict:
    """One dashboard sheet. `cells` decides what a cell holds — the make path's
    plain values, or the render path's links, notes and citations."""
    ws.sheet_view.showGridLines = False       # gridlines are the "spreadsheet" tell

    heads = [cells.head(h) for h in (spec.get("headers") or [])]
    headers = [cells.shown(h) for h in heads]
    rows = spec.get("rows") or []
    ncols = max(len(headers), 4)
    last_col = ncols

    # ── title band ───────────────────────────────────────────────────────────
    _band(ws, 1, 1, last_col, t, t.deep, height=34)
    ws.cell(row=1, column=1, value=cells.text(spec.get("title", ""))).font = _font(
        t, size=16, bold=True, colour=t.on(t.deep))
    ws.cell(row=1, column=1).alignment = Alignment(vertical="center", indent=1)
    if spec.get("subtitle"):
        _band(ws, 2, 1, last_col, t, t.deep, height=20)
        c = ws.cell(row=2, column=1, value=cells.text(spec["subtitle"]))
        c.font = _font(t, size=10, colour=t.support)
        c.alignment = Alignment(vertical="center", indent=1)
    row = 4

    # ── KPI row: merged cells, oversized numerals ────────────────────────────
    kpis = (spec.get("kpis") or [])[:4]
    if kpis:
        span = max(1, ncols // len(kpis))
        for i, k in enumerate(kpis):
            c0 = 1 + i * span
            c1 = min(last_col, c0 + span - 1)
            ws.merge_cells(start_row=row, start_column=c0, end_row=row, end_column=c1)
            cell = ws.cell(row=row, column=c0, value=cells.text(k.get("value")))
            cell.font = _font(t, size=22, bold=True, colour=t.primary)
            # A KPI written as a bare number renders as 1285000, which is the
            # exact "data dump" look this renderer exists to avoid. The model
            # may name a format per KPI; otherwise a numeric value gets grouped.
            if isinstance(k.get("value"), (int, float)):
                cell.number_format = k.get("format") or spec.get("kpi_format") or "#,##0"
            cell.alignment = Alignment(vertical="center", indent=1)
            ws.merge_cells(start_row=row + 1, start_column=c0, end_row=row + 1, end_column=c1)
            lab = ws.cell(row=row + 1, column=c0, value=cells.text(k.get("label", "")))
            lab.font = _font(t, size=9, colour=t.mute)
            lab.alignment = Alignment(vertical="top", indent=1)
        ws.row_dimensions[row].height = 30
        ws.row_dimensions[row + 1].height = 16
        row += 3

    # ── data table ───────────────────────────────────────────────────────────
    head_row = row
    thin = Side(style="thin", color=_hx(t.faint))
    for c, h in enumerate(heads, start=1):
        cell = cells.put(ws, head_row, c, h)
        cell.fill = _fill(t.primary)
        cell.font = _font(t, size=10, bold=True, colour=t.on(t.primary))
        cell.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[head_row].height = 22

    numeric_cols: set[int] = set()
    for r, data in enumerate(rows, start=head_row + 1):
        for c in range(1, len(headers) + 1):
            v = data[c - 1] if isinstance(data, list) and c <= len(data) else None
            # A string starting with '=' is a FORMULA and must stay one.
            cell = cells.put(ws, r, c, v)
            cell.font = _font(t, size=10, bold=(c == 1))
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="center", indent=1 if c == 1 else 0,
                                       horizontal="left" if c == 1 else "right")
            if isinstance(cell.value, (int, float)):
                numeric_cols.add(c)
                cell.number_format = spec.get("number_format", "#,##0")
            if (r - head_row) % 2 == 0:
                cell.fill = _fill(t.faint)
            cells.style(cell, t)
    last_row = head_row + len(rows)

    # ── a data bar on the main numeric column: proportion at a glance ────────
    if numeric_cols and last_row > head_row:
        col = get_column_letter(min(numeric_cols))
        ws.conditional_formatting.add(
            f"{col}{head_row+1}:{col}{last_row}",
            DataBarRule(start_type="min", end_type="max", color=_hx(t.accent),
                        showValue=True))

    # ── totals row, as a real formula ───────────────────────────────────────
    if spec.get("total_row") and numeric_cols and last_row > head_row:
        tr = last_row + 1
        ws.cell(row=tr, column=1, value=cells.text(spec.get("total_label", "Total"))).font = _font(
            t, size=10, bold=True, colour=t.on(t.deep))
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=tr, column=c)
            cell.fill = _fill(t.deep)
            cell.font = _font(t, size=10, bold=True, colour=t.on(t.deep))
            if c in numeric_cols:
                col = get_column_letter(c)
                cell.value = f"=SUM({col}{head_row+1}:{col}{last_row})"
                cell.number_format = spec.get("number_format", "#,##0")
                cell.alignment = Alignment(horizontal="right")
        ws.row_dimensions[tr].height = 20
        last_row = tr

    # ── native chart, bound to the cells ────────────────────────────────────
    chart_kind = (spec.get("chart") or {}).get("type")
    if chart_kind and numeric_cols and len(rows) >= 2:
        ch = {"bar": BarChart, "line": LineChart, "pie": PieChart}.get(chart_kind, BarChart)()
        vcol = min(numeric_cols)
        data = Reference(ws, min_col=vcol, min_row=head_row,
                         max_row=head_row + len(rows))
        cats = Reference(ws, min_col=1, min_row=head_row + 1,
                         max_row=head_row + len(rows))
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
        ch.title = cells.text((spec.get("chart") or {}).get("title")) or None
        ch.height, ch.width = 7.5, 15
        if hasattr(ch, "gapWidth"):
            ch.gapWidth = 60
        ws.add_chart(ch, f"A{last_row + 3}")

    # ── the mechanical polish nobody should spend tokens on ─────────────────
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    for c in range(1, len(headers) + 1):
        longest = len(cells.shown(headers[c - 1]))
        for data in rows:
            if isinstance(data, list) and c <= len(data):
                longest = max(longest, len(cells.shown(data[c - 1])))
        ws.column_dimensions[get_column_letter(c)].width = min(46, max(12, longest + 6))
    if headers:
        ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(headers))}{head_row + len(rows)}"
    return {"sheet": ws.title, "kind": "table", "title": cells.shown(spec.get("title", "")),
            "headers": [cells.shown(h) for h in headers], "rows": len(rows),
            "chart": chart_kind or None}


def build(spec: dict, out: Path, *, drawn: list | None = None, warnings: list | None = None,
          cite=None) -> Path:
    """Render the workbook. `cite` (citations.Citations) is the render path's:
    given, the spec may hold several sheets, rich cells and sources; make
    never passes it and gets its one sheet exactly as before."""
    t = pal.from_spec(spec)
    wb = Workbook()
    if cite is None:
        ws = wb.active
        ws.title = (spec.get("sheet_name") or "Dashboard")[:31]
        _sheet(ws, spec, t)
    else:
        _workbook(wb, spec, t, cite, drawn if drawn is not None else [],
                  warnings if warnings is not None else [])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    return out


# ── the render path (WF-06) ──────────────────────────────────────────────────
_BAD_NAME = re.compile(r"[\[\]:*?/\\]")
SOURCES_HEAD_ROW = 4          # the Sources sheet's header row; source n sits at row 4 + n


def _sheet_name(raw, taken: set[str], fallback: str) -> str:
    name = _BAD_NAME.sub("-", " ".join(str(raw or "").split())).strip("'") or fallback
    name = name[:31]
    base, k = name, 2
    while name.lower() in taken:
        suffix = f" ({k})"
        name = base[:31 - len(suffix)] + suffix
        k += 1
    taken.add(name.lower())
    return name


class _Rich:
    """How the render path writes a cell: a value or {"text"|"value", "url",
    "note", "cite"}; citation markers in text as "[3]"."""

    def __init__(self, cite, sources_sheet: str | None):
        self.cite = cite
        self.sources_sheet = sources_sheet
        self.links = 0
        self.notes = 0

    @staticmethod
    def head(h):
        return h                       # a header may be a rich cell too (a note)

    def text(self, v):
        if isinstance(v, str):
            return self.cite.plain(v)
        if isinstance(v, dict):
            return self.text(v.get("text", v.get("value")))
        return v

    def shown(self, v) -> str:
        if isinstance(v, dict) and v.get("cite") is not None and v.get("text") is None and v.get("value") is None:
            return self._cite_text(v)
        return str(self.text(v) if self.text(v) is not None else "")

    def _cite_text(self, v: dict) -> str:
        from citations import numbers_text
        return numbers_text(self.cite.resolve_ids(v.get("cite")))

    def put(self, ws, row, col, v):
        cell = ws.cell(row=row, column=col)
        if not isinstance(v, dict):
            cell.value = self.text(v)
            return cell
        link = None
        if v.get("cite") is not None:
            nums = self.cite.resolve_ids(v.get("cite"))
            from citations import numbers_text
            cell.value = self.text(v) if (v.get("text") is not None or v.get("value") is not None) \
                else numbers_text(nums)
            if nums and self.sources_sheet:
                link = ("internal", f"'{self.sources_sheet}'!A{SOURCES_HEAD_ROW + min(nums)}")
        else:
            cell.value = self.text(v)
        url = str(v.get("url") or "").strip()
        if url.startswith("#"):
            link = ("internal", url[1:])
        elif re.match(r"^https?://\S+$", url, re.I):
            link = ("external", url)
        elif url:
            self.cite.warnings.append(f"cell {cell.coordinate} on '{ws.title}': '{url[:60]}' is not an "
                                      "http(s) URL; shown without a link")
        if link is not None:
            if link[0] == "external":
                cell.hyperlink = link[1]
            else:
                cell.hyperlink = Hyperlink(ref=cell.coordinate, location=link[1])
            self.links += 1
        note = self.cite.plain(" ".join(str(v.get("note") or "").split()))
        if note:
            c = Comment(note, "Bobble")
            c.width, c.height = 260, max(60, 18 * (1 + len(note) // 38))
            cell.comment = c
            self.notes += 1
        return cell

    def style(self, cell, t):
        """A linked cell reads as a link: the theme's primary, underlined."""
        if cell.hyperlink is not None:
            f = cell.font
            cell.font = Font(name=f.name, size=f.sz, bold=f.b, italic=f.i, underline="single",
                             color=_hx(t.primary))


def _sources_sheet(ws, spec: dict, t, cite, cells: _Rich) -> dict:
    """Every source, one row each, in the citations' numbering; the URL a live
    link. Columns nobody filled are left out rather than shown empty."""
    ws.sheet_view.showGridLines = False
    cols = [("#", lambda s: s.n), ("Title", lambda s: s.title)]
    for head, get in (("Site", lambda s: s.site), ("Date", lambda s: s.date)):
        if any(get(s) for s in cite.sources):
            cols.append((head, get))
    cols.append(("URL", lambda s: s.url))
    if any(s.accessed for s in cite.sources):
        cols.append(("Accessed", lambda s: s.accessed))
    last_col = max(len(cols), 4)
    _band(ws, 1, 1, last_col, t, t.deep, height=34)
    title = ws.cell(row=1, column=1, value=cells.text(spec.get("title")) or "Sources")
    title.font = _font(t, size=16, bold=True, colour=t.on(t.deep))
    title.alignment = Alignment(vertical="center", indent=1)
    if spec.get("subtitle"):
        _band(ws, 2, 1, last_col, t, t.deep, height=20)
        c = ws.cell(row=2, column=1, value=cells.text(spec["subtitle"]))
        c.font = _font(t, size=10, colour=t.support)
        c.alignment = Alignment(vertical="center", indent=1)
    thin = Side(style="thin", color=_hx(t.faint))
    for c, (head, _g) in enumerate(cols, start=1):
        cell = ws.cell(row=SOURCES_HEAD_ROW, column=c, value=head)
        cell.fill = _fill(t.primary)
        cell.font = _font(t, size=10, bold=True, colour=t.on(t.primary))
        cell.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[SOURCES_HEAD_ROW].height = 22
    widths = [len(h) for h, _g in cols]
    for src in cite.sources:
        r = SOURCES_HEAD_ROW + src.n
        for c, (head, get) in enumerate(cols, start=1):
            v = get(src)
            cell = ws.cell(row=r, column=c, value=v if v != "" else None)
            cell.font = _font(t, size=10, bold=(head == "Title"))
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="center", indent=1 if c == 1 else 0,
                                       horizontal="left", wrap_text=head == "Title")
            if head == "URL" and v:
                cell.hyperlink = v
                cell.font = Font(name=FONT, size=10, underline="single", color=_hx(t.primary))
                cells.links += 1
            if src.n % 2 == 0:
                cell.fill = _fill(t.faint)
            widths[c - 1] = max(widths[c - 1], len(str(v)))
    for c, w in enumerate(widths, start=1):
        cap = 6 if c == 1 else 48 if cols[c - 1][0] == "Title" else 70
        ws.column_dimensions[get_column_letter(c)].width = min(cap, max(6 if c == 1 else 12, w + 4))
    ws.freeze_panes = ws.cell(row=SOURCES_HEAD_ROW + 1, column=1)
    return {"sheet": ws.title, "kind": "sources", "title": title.value, "headers": [h for h, _g in cols],
            "rows": len(cite.sources), "links": sum(1 for s in cite.sources if s.url)}


def _workbook(wb, spec: dict, t, cite, drawn: list, warnings: list) -> None:
    parts = spec.get("sheets") if isinstance(spec.get("sheets"), list) and spec.get("sheets") else [spec]
    parts = [p for p in parts if isinstance(p, dict)]
    has_sources_part = any(p.get("kind") == "sources" for p in parts)
    if cite.sources and not has_sources_part:
        parts = parts + [{"kind": "sources", "name": "Sources"}]
    taken: set[str] = set()
    names = [_sheet_name(p.get("name") or p.get("sheet_name"),
                         taken, "Sources" if p.get("kind") == "sources" else f"Sheet{i}")
             for i, p in enumerate(parts, 1)]
    src_name = next((n for n, p in zip(names, parts) if p.get("kind") == "sources"), None)
    for i, (part, name) in enumerate(zip(parts, names)):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = name
        cells = _Rich(cite, src_name if cite.sources else None)
        if part.get("kind") == "sources":
            if not cite.sources:
                warnings.append(f"sheet '{name}': the spec has no sources to list")
            rec = _sources_sheet(ws, part, t, cite, cells)
        else:
            rec = _sheet(ws, part, t, cells)
            rec["links"] = cells.links
        rec["notes"] = cells.notes
        drawn.append(rec)


if __name__ == "__main__":
    spec = json.loads(Path(sys.argv[1]).read_text())
    p = build(spec, Path(sys.argv[2]))
    print(f"wrote {p}")
