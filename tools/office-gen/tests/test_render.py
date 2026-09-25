"""
WF-06 — `office.py render`: a spec drawn exactly, with its sources.

The render path is the one a research workflow calls with a spec it built from
pages it fetched (deliverables/research/workflows.md §4.5). What it must do:
  - make NO model call and invent nothing;
  - docx: citation superscripts that link to a Sources block whose URLs are
    live links, and a contents list that links to its headings;
  - pptx: "Sources: 3, 7" lines and a sources slide with live links;
  - xlsx: several sheets, cells with links and notes, a Sources sheet;
  - `office inspect` shows the sources with their URLs;
and none of it may leak into `make`, whose specs come from a small model.

    tools/office-gen/.venv/bin/python -m pytest tools/office-gen/tests/test_render.py -q
"""
from __future__ import annotations

import json
import os
import re
import socket
import subprocess
import sys
import time
import zipfile
from pathlib import Path

import pytest

HERE = Path(__file__).resolve().parent
OFFICE = HERE.parent
REPO = OFFICE.parents[1]
EVAL = REPO / "tools" / "visual-eval"
RENDER_FIX = HERE / "fixtures" / "render"
sys.path.insert(0, str(OFFICE))

import citations  # noqa: E402
import doc_render  # noqa: E402
import office  # noqa: E402
import office_edit  # noqa: E402
import render_deck  # noqa: E402

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def fixture(name: str) -> dict:
    return json.loads((RENDER_FIX / f"{name}.json").read_text())


def urls(spec: dict) -> set[str]:
    return {s["url"] for s in spec["sources"]}


@pytest.fixture(autouse=True)
def _scratch(tmp_path, monkeypatch):
    monkeypatch.setenv("PI_OFFICE_GEN_SCRATCH", str(tmp_path / "scratch"))


# ── the command, as a subprocess, with a model server that records calls ─────
@pytest.fixture(scope="module")
def recorder(tmp_path_factory):
    """A model server that answers nothing and writes down every request: the
    render path must leave its log empty."""
    d = tmp_path_factory.mktemp("recorder")
    replies = d / "replies.json"
    replies.write_text("{}")
    log = d / "requests.jsonl"
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    srv = subprocess.Popen([sys.executable, str(EVAL / "py" / "replay_server.py"), str(replies), str(port), str(log)])
    for _ in range(100):
        try:
            socket.create_connection(("127.0.0.1", port), timeout=0.2).close()
            break
        except OSError:
            time.sleep(0.05)
    yield f"http://127.0.0.1:{port}", log
    srv.terminate()
    srv.wait(timeout=5)


def run_office(args: list[str], server: str, tmp_path: Path) -> dict:
    env = {**os.environ, "PI_OFFICE_GEN_SERVER": server, "PI_OFFICE_GEN_SCRATCH": str(tmp_path / "s"),
           "no_proxy": "*", "NO_PROXY": "*"}
    r = subprocess.run([sys.executable, str(OFFICE / "office.py"), *args], env=env, capture_output=True,
                       text=True, timeout=120)
    lines = r.stdout.strip().splitlines()
    assert len(lines) == 1, r.stdout + r.stderr          # one JSON object on stdout, nothing else
    return json.loads(lines[0])


@pytest.mark.parametrize("kind,name", [("docx", "report"), ("pptx", "deck"), ("xlsx", "workbook"),
                                       ("pdf", "report")])
def test_render_makes_no_model_call(recorder, tmp_path, kind, name):
    server, log = recorder
    out = run_office(["render", kind, "--spec", str(RENDER_FIX / f"{name}.json"),
                      "--out", str(tmp_path / f"{name}.{kind}")], server, tmp_path)
    assert out["ok"], out
    assert out["sources"] == 5 and out["checks"]["unknown_citations"] == []
    assert not log.exists() or log.read_text() == "", "render asked the model server something"


def test_render_chart_too(recorder, tmp_path):
    server, log = recorder
    spec = json.loads((EVAL / "fixtures" / "captured-4b" / "last_chart_spec.json").read_text())
    f = tmp_path / "chart.json"
    f.write_text(json.dumps(spec))
    out = run_office(["render", "chart", "--spec", str(f), "--out", str(tmp_path / "units.svg")], server, tmp_path)
    assert out["ok"], out
    assert (tmp_path / "units.chart.json").exists()        # `office edit units.svg` revises this
    assert not log.exists() or log.read_text() == ""


def test_a_spec_of_the_wrong_shape_fails_loudly(recorder, tmp_path):
    server, _log = recorder
    f = tmp_path / "bad.json"
    f.write_text(json.dumps({"slides": []}))
    out = run_office(["render", "pptx", "--spec", str(f), "--out", str(tmp_path / "x.pptx")], server, tmp_path)
    assert out["ok"] is False and "slides" in out["error"]
    out = run_office(["render", "docx", "--spec", str(tmp_path / "missing.json"), "--out",
                      str(tmp_path / "x.docx")], server, tmp_path)
    assert out["ok"] is False and "no spec file" in out["error"]


def test_make_takes_its_brief_from_a_file(tmp_path):
    """--brief-file: the brief a workflow writes to disk reaches the model whole."""
    rep = {"doc": {"title": "Q3 support", "palette": {"primary": "#1F3A5F", "accent": "#C8742B"},
                   "blocks": [{"type": "heading", "title": "Q3 support"},
                              {"type": "body", "paragraphs": ["12,480 tickets, up 9% on Q2."]}]}}
    rfile = tmp_path / "replies.json"
    rfile.write_text(json.dumps(rep))
    log = tmp_path / "requests.jsonl"
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    srv = subprocess.Popen([sys.executable, str(EVAL / "py" / "replay_server.py"), str(rfile), str(port), str(log)])
    try:
        for _ in range(100):
            try:
                socket.create_connection(("127.0.0.1", port), timeout=0.2).close()
                break
            except OSError:
                time.sleep(0.05)
        brief = "A one-page report on Q3 support: 12,480 tickets, up 9% on Q2.\n" + "Context line. " * 400
        bf = tmp_path / "brief.txt"
        bf.write_text(brief)
        out = run_office(["make", "docx", "--brief-file", str(bf), "--out", str(tmp_path / "r.docx")],
                         f"http://127.0.0.1:{port}", tmp_path)
        both = run_office(["make", "docx", "--brief", "x" * 20, "--brief-file", str(bf), "--out",
                           str(tmp_path / "r2.docx")], f"http://127.0.0.1:{port}", tmp_path)
        neither = run_office(["make", "docx", "--out", str(tmp_path / "r3.docx")], f"http://127.0.0.1:{port}",
                             tmp_path)
    finally:
        srv.terminate()
        srv.wait(timeout=5)
    assert out["ok"], out
    sent = [json.loads(ln) for ln in log.read_text().splitlines()]
    assert brief.strip() in sent[0]["user"]
    assert both["ok"] is False and "not both" in both["error"]
    assert neither["ok"] is False and "--brief-file" in neither["error"]


# ── docx ─────────────────────────────────────────────────────────────────────
@pytest.fixture(scope="module")
def report(tmp_path_factory):
    d = tmp_path_factory.mktemp("report")
    spec = fixture("report")
    out = office.render("docx", spec, d / "report.docx")
    return spec, out, d / "report.docx"


def _docx_xml(path: Path):
    from lxml import etree
    with zipfile.ZipFile(path) as z:
        doc = etree.fromstring(z.read("word/document.xml"))
        rels = etree.fromstring(z.read("word/_rels/document.xml.rels"))
    return doc, rels


def test_docx_sources_are_live_links_and_nothing_else_is(report):
    spec, out, path = report
    doc, rels = _docx_xml(path)
    external = {r.get("Id"): r.get("Target") for r in rels
                if r.get("Type", "").endswith("/hyperlink") and r.get("TargetMode") == "External"}
    linked = [external[h.get(f"{R}id")] for h in doc.iter(f"{W}hyperlink") if h.get(f"{R}id")]
    assert sorted(linked) == sorted(urls(spec))            # one link per source, each its own URL
    assert set(external.values()) == urls(spec)            # no URL the spec did not give


def test_docx_citations_link_to_their_entries(report):
    spec, out, path = report
    doc, _rels = _docx_xml(path)
    bookmarks = {b.get(f"{W}name") for b in doc.iter(f"{W}bookmarkStart")}
    anchors = [h.get(f"{W}anchor") for h in doc.iter(f"{W}hyperlink") if h.get(f"{W}anchor")]
    cited = [a for a in anchors if a.startswith("_RefSource")]
    assert cited and all(a in bookmarks for a in anchors)  # every internal link lands somewhere
    # One superscript link per id the markers name, in the order written; a
    # run of markers ("[S4][S2]") is one group, numbered in order (²˒⁴).
    groups = re.findall(r"(?:\[\s*S\d+(?:\s*[,;]\s*S\d+)*\s*\])+", json.dumps(spec["blocks"]))
    marked = [str(n) for g in groups for n in sorted({int(x) for x in re.findall(r"S(\d+)", g)})]
    assert [a.removeprefix("_RefSource") for a in cited] == marked
    for h in doc.iter(f"{W}hyperlink"):
        if (h.get(f"{W}anchor") or "").startswith("_RefSource"):
            assert h.find(f".//{W}vertAlign").get(f"{W}val") == "superscript"
    body_text = "".join(t.text or "" for t in doc.iter(f"{W}t"))
    assert "[S" not in body_text


def test_docx_contents_link_to_the_headings(report):
    spec, out, path = report
    doc, _rels = _docx_xml(path)
    toc = [h for h in doc.iter(f"{W}hyperlink") if (h.get(f"{W}anchor") or "").startswith("_TocSection")]
    titles = ["".join(t.text for t in h.iter(f"{W}t")) for h in toc]
    assert titles == ["Summary", "Findings", "What we could not confirm", "Sources"]
    for h in toc:
        target = h.get(f"{W}anchor")
        para = next(b.getparent() for b in doc.iter(f"{W}bookmarkStart") if b.get(f"{W}name") == target)
        assert "".join(t.text or "" for t in para.iter(f"{W}t")) in titles


def test_docx_invents_nothing(report):
    """Every word in the file is the spec's, or one of the renderer's three
    labels; every number is the spec's or a source's number."""
    spec, out, path = report
    doc, _rels = _docx_xml(path)
    text = " ".join(t.text or "" for t in doc.iter(f"{W}t"))
    allowed = set(re.findall(r"[a-z]+", json.dumps(spec).lower())) | {"sources", "contents", "accessed"}
    assert set(re.findall(r"[a-z]+", text.lower())) <= allowed
    spec_numbers = set(re.findall(r"\d[\d,.]*", json.dumps(spec)))
    ns = {str(i) for i in range(1, len(spec["sources"]) + 1)}
    assert set(re.findall(r"\d[\d,.]*", text)) - spec_numbers - ns == set()


def test_docx_inspect_shows_the_sources_with_their_urls(report, tmp_path):
    spec, out, path = report
    outline = office.inspect(path)["outline"]
    for u in urls(spec):
        assert f"→ {u}" in outline
    # A superscript read flat runs into its word ("claim2", "71%3"); the
    # outline says it as a citation.
    assert "by 71%[3]" in outline and "insurance claim[2]." in outline and "buildings[2,4]." in outline
    assert 't2.r1.c2:"Insurance data[2]"' in outline
    assert '"1 Tidewell pilot results, Q3 2026 — Tidewell, 12 Sep 2026 https://' in outline


def test_a_long_documents_sources_survive_the_outline_cut(tmp_path):
    spec = fixture("report")
    spec["blocks"] = spec["blocks"][:4] + [{"type": "body", "paragraphs": [f"Paragraph {i} [S1]."]}
                                           for i in range(100)] + [{"type": "sources"}]
    path = tmp_path / "long.docx"
    office.render("docx", spec, path)
    outline = office.inspect(path)["outline"]
    assert len(outline.splitlines()) > office.OUTLINE_LINES
    assert "with links follow" in outline
    for u in urls(spec):
        assert f"→ {u}" in outline


def test_editing_a_cited_paragraph_leaves_no_old_link_behind(report, tmp_path):
    spec, out, path = report
    o = office_edit.inspect_docx(path)
    pid = next(p["id"] for p in o["paragraphs"] if p["text"].startswith("Water damage is the most common"))
    dst = tmp_path / "edited.docx"
    rep = office_edit.apply_docx(path, [{"op": "set_text", "id": pid, "text": "Reworded."}], dst)
    assert rep == [f"ok set_text {pid}"]
    edited = next(p for p in office_edit.inspect_docx(dst)["paragraphs"] if p["id"] == pid)
    assert edited["text"] == "Reworded."


# ── pptx ─────────────────────────────────────────────────────────────────────
@pytest.fixture(scope="module")
def deck(tmp_path_factory):
    d = tmp_path_factory.mktemp("deck")
    spec = fixture("deck")
    out = office.render("pptx", spec, d / "deck.pptx")
    from pptx import Presentation
    return spec, out, Presentation(str(d / "deck.pptx")), d / "deck.pptx"


def _texts(slide):
    return [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]


def test_pptx_sources_slide_lists_every_source_as_a_live_link(deck):
    spec, out, prs, _path = deck
    last = list(prs.slides)[-1]
    links = [r.hyperlink.address for sh in _texts(last) for p in sh.text_frame.paragraphs for r in p.runs
             if r.hyperlink.address]
    assert sorted(links) == sorted(urls(spec))
    assert out["summary"].splitlines()[-1].startswith("6. sources: Sources")


def test_pptx_each_cited_slide_says_its_sources_where_nothing_else_is(deck):
    spec, out, prs, _path = deck
    want = {2: "Sources: 2–4", 3: "Sources: 1–3, 5", 4: "Sources: 1, 5", 5: "Sources: 3"}
    slides = list(prs.slides)
    for n, text in want.items():
        sl = slides[n - 1]
        lines = [sh for sh in _texts(sl) if sh.text_frame.text.startswith("Sources:")]
        assert [sh.text_frame.text for sh in lines] == [text], n
        line = lines[0]
        # …clear of every other piece of text on the slide…
        for sh in _texts(sl):
            if sh.shape_id == line.shape_id:
                continue
            assert (line.left + line.width <= sh.left or sh.left + sh.width <= line.left or
                    line.top + line.height <= sh.top or sh.top + sh.height <= line.top), (n, sh.text_frame.text)
        # …and a link to the slide that lists the first of them.
        assert line.click_action.target_slide == slides[-1]
    assert not any("[S" in sh.text_frame.text for sl in slides for sh in _texts(sl))
    assert not [sh for sh in _texts(slides[0]) if sh.text_frame.text.startswith("Sources:")]


def test_a_long_source_list_continues_onto_more_slides(tmp_path):
    spec = fixture("deck")
    base = spec["sources"]
    spec["sources"] = [{**base[i % 5], "id": f"S{i + 1}", "title": f"{base[i % 5]['title']} (part {i + 1})"}
                       for i in range(14)]
    out = office.render("pptx", spec, tmp_path / "long.pptx")
    from pptx import Presentation
    slides = list(Presentation(str(tmp_path / "long.pptx")).slides)
    pages = [sl for sl in slides if any(sh.text_frame.text.startswith("Sources") and sh.top < 1600000
                                        for sh in _texts(sl))]
    assert len(pages) >= 2
    listed = [r.hyperlink.address for sl in pages for sh in _texts(sl) for p in sh.text_frame.paragraphs
              for r in p.runs if r.hyperlink.address]
    assert len(listed) == 14                               # every source, once — never cut
    for sl in pages:
        for sh in _texts(sl):
            assert sh.top + sh.height <= render_deck.H, "an entry runs off the slide"
    assert out["ok"] and not any("shown" in w for w in out["warnings"])


@pytest.mark.parametrize("title", ["Sources", "Sources, continued",
                                   "Where every figure in this deck came from, source by source"])
def test_planned_head_is_the_drawn_head(title):
    from pptx import Presentation
    prs = Presentation()
    prs.slide_width, prs.slide_height = render_deck.W, render_deck.H
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    drawn = render_deck._head(sl, render_deck.THEMES["ink"], {"title": title})
    assert drawn == render_deck._head_bottom(title)


# ── xlsx ─────────────────────────────────────────────────────────────────────
@pytest.fixture(scope="module")
def workbook(tmp_path_factory):
    d = tmp_path_factory.mktemp("workbook")
    spec = fixture("workbook")
    out = office.render("xlsx", spec, d / "workbook.xlsx")
    from openpyxl import load_workbook
    return spec, out, load_workbook(d / "workbook.xlsx"), d / "workbook.xlsx"


def test_xlsx_has_its_sheets_and_a_sources_sheet_of_live_links(workbook):
    spec, out, wb, _path = workbook
    assert wb.sheetnames == ["Findings", "Data", "Sources"]
    ws = wb["Sources"]
    linked = {c.hyperlink.target for row in ws.iter_rows() for c in row if c.hyperlink is not None}
    assert linked == urls(spec)
    for src_n in range(1, 6):
        assert ws.cell(row=4 + src_n, column=1).value == src_n
    assert len(wb["Data"]._charts) == 1


def test_xlsx_cite_cells_link_to_their_source_rows(workbook):
    spec, out, wb, _path = workbook
    ws = wb["Findings"]
    cites = {c.value: c.hyperlink.location for row in ws.iter_rows() for c in row
             if c.hyperlink is not None and c.hyperlink.location}
    # One source is its number (3, not "3": text that looks like a number is
    # flagged as an error by a spreadsheet); two are "2, 4".
    assert cites == {"2, 4": "'Sources'!A6", 3: "'Sources'!A7", 5: "'Sources'!A9", 1: "'Sources'!A5"}
    src = wb["Sources"]
    for loc in cites.values():
        row = int(loc.rsplit("A", 1)[1])
        assert isinstance(src.cell(row=row, column=1).value, int)
    # A citation's number is not data: no data bar over it (SEEN in the canvas
    # before this), no thousands format, set left like the text around it.
    assert [str(cf.sqref) for cf in ws.conditional_formatting] == []
    assert ws["C6"].number_format == "General" and ws["C6"].alignment.horizontal == "left"
    assert [str(cf.sqref) for cf in wb["Data"].conditional_formatting] == ["B5:B8"]


def test_xlsx_sentences_wrap_and_their_rows_grow(workbook):
    """A claim is a sentence: clipped at the next cell it was not readable
    (SEEN in the canvas). It wraps, left-aligned, in a row tall enough."""
    spec, out, wb, _path = workbook
    ws = wb["Findings"]
    claim = ws["A5"]
    assert claim.alignment.wrap_text and claim.alignment.horizontal == "left"
    assert ws.row_dimensions[5].height and ws.row_dimensions[5].height > 20
    assert ws["B5"].alignment.horizontal == "left"            # text, not a number
    src = wb["Sources"]
    assert [c.value for c in src[4]][:3] == ["#", "Title", "URL"]


def test_xlsx_notes_and_no_leftover_markers(workbook):
    spec, out, wb, _path = workbook
    notes = {(ws.title, c.coordinate): c.comment.text for ws in wb.worksheets for row in ws.iter_rows()
             for c in row if c.comment is not None}
    assert notes == {("Findings", "B7"): "One vendor benchmark [5]; no independent measurement found.",
                     ("Data", "A5"): "The #1 cause in every year of the series [2]."}
    assert not any("[S" in str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
                   if c.value is not None)


def test_xlsx_inspect_shows_links_and_notes(workbook):
    spec, out, wb, path = workbook
    outline = office.inspect(path)["outline"]
    for u in urls(spec):
        assert f"→ {u}" in outline
    assert "→ #'Sources'!A6" in outline
    assert 'note: "One vendor benchmark [5]' in outline


# ── pdf ──────────────────────────────────────────────────────────────────────
def test_pdf_sources_are_link_annotations(tmp_path):
    spec = fixture("report")
    out = office.render("pdf", spec, tmp_path / "report.pdf")
    import pypdf
    reader = pypdf.PdfReader(str(tmp_path / "report.pdf"))
    uris = {a.get_object()["/A"]["/URI"] for p in reader.pages for a in (p.get("/Annots") or [])
            if a.get_object().get("/A", {}).get("/URI")}
    assert uris == urls(spec)
    text = " ".join(p.extract_text() for p in reader.pages)
    assert "[S" not in text and "[2]" in text
    # The contents list links to its headings, and the viewer's outline has them.
    assert [o.title for o in reader.outline] == ["Summary", "Findings", "What we could not confirm", "Sources"]
    internal = [a.get_object() for p in reader.pages for a in (p.get("/Annots") or [])
                if a.get_object().get("/Dest") is not None]
    assert len(internal) == 4


@pytest.mark.parametrize("which", ["report", "deck"])
def test_no_followed_link_purple(which, report, deck):
    """The templates' followed-link colour is 800080, which PowerPoint paints a
    clicked link in whatever the run says; a rendered file's links stay the
    document's own colour (the user: no purple, anywhere)."""
    path = report[2] if which == "report" else deck[3]
    part = "word/theme/theme1.xml" if which == "report" else "ppt/theme/theme1.xml"
    with zipfile.ZipFile(path) as z:
        theme = z.read(part).decode()
    link = re.search(r"<a:hlink>\s*<a:srgbClr val=\"([0-9A-F]{6})\"/>\s*</a:hlink>", theme).group(1)
    followed = re.search(r"<a:folHlink>\s*<a:srgbClr val=\"([0-9A-F]{6})\"/>\s*</a:folHlink>", theme).group(1)
    assert link == followed != "800080"
    assert link == ("1F3A5F" if which == "report" else str(render_deck.THEMES["ink"].primary))


# ── the make path is untouched ───────────────────────────────────────────────
def test_make_never_draws_a_bibliography_a_model_wrote(tmp_path):
    """Without `cite` (every make) a sources block is an unknown block set as
    body text, a marker is just text, and the file has no link of any kind."""
    spec = {"title": "t", "palette": {"primary": "#1F3A5F", "accent": "#C8742B"},
            "sources": [{"id": "S1", "title": "Invented", "url": "https://invented.example/x"}],
            "blocks": [{"type": "body", "paragraphs": ["A claim [S1]."]},
                       {"type": "toc"}, {"type": "sources", "text": "a list"}]}
    warn: list[str] = []
    doc_render.build(spec, tmp_path / "m.docx", warnings=warn)
    doc, rels = _docx_xml(tmp_path / "m.docx")
    assert not list(doc.iter(f"{W}hyperlink")) and not list(doc.iter(f"{W}bookmarkStart"))
    assert not [r for r in rels if r.get("Type", "").endswith("/hyperlink")]
    assert "A claim [S1]." in "".join(t.text or "" for t in doc.iter(f"{W}t"))
    assert any("'sources' does not exist" in w for w in warn)
    deck_spec = {"slides": [{"layout": "bullets", "title": "x", "bullets": ["b [S1]"], "sources": ["S1"]},
                            {"layout": "sources"}], "sources": spec["sources"]}
    render_deck.build(deck_spec, tmp_path / "m.pptx")
    from pptx import Presentation
    prs = Presentation(str(tmp_path / "m.pptx"))
    assert len(prs.slides) == 2
    assert not any(sh.text_frame.text.startswith("Sources:") for sl in prs.slides for sh in _texts(sl))


# ── citations.py ─────────────────────────────────────────────────────────────
def test_markers_number_by_position_and_unknown_ids_are_reported():
    c = citations.Citations.of([{"id": "S3", "title": "a", "url": "https://a.example"},
                                {"id": "S7", "title": "b", "url": "https://b.example"}])
    assert c.split("rose 14% [S7][S3].") == ["rose 14%", [1, 2], "."]
    assert c.strip("x [S3; S7] y") == ("x y", [1, 2])
    assert c.plain("x [s7] y") == "x [2] y"
    assert c.split("[1] and [sic] stay") == ["[1] and [sic] stay"]
    assert c.split("gone [S9].") == ["gone", "."] and c.unknown == ["S9"]
    assert "name no source" in c.report()[-1]


def test_numbers_as_ranges():
    assert citations.numbers_text([7, 1, 2, 3]) == "1–3, 7"
    assert citations.numbers_text([1, 2]) == "1, 2"
    assert citations.numbers_text([4]) == "4"


def test_a_bad_source_is_dropped_or_unlinked_and_said():
    c = citations.Citations.of([{"title": "no url"}, {"url": "ftp://x"}, {"id": "S1", "url": "https://ok.example"},
                                {"id": "s1", "url": "https://dup.example"}, "junk", {}])
    assert [(s.id, s.title) for s in c.sources] == [("#1", "no url"), ("S1", "https://ok.example")]
    assert c.sources[0].url == ""
    said = " | ".join(c.warnings)
    for why in ("not an http(s) URL", "neither a URL nor a title", "used twice", "not an object"):
        assert why in said
    assert c.split("see [S1]") == ["see", [2]]          # the id that was claimed, not the entry without one
