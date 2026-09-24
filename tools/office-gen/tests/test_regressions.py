"""
VQ-01 — the office renderers against the defects the visual-quality research
reproduced (deliverables/research/visual-quality.md §2.3), on the REAL 4B
captures and the research's replays of 4B habits (tools/visual-eval/fixtures).

Geometry is judged by the eval's own ruler (tools/visual-eval/py/measure.py),
which reads the written FILE, not the renderer's intentions.

    tools/office-gen/.venv/bin/python -m pytest tools/office-gen/tests -q
"""
from __future__ import annotations

import json
import os
import socket
import subprocess
import sys
import time
from pathlib import Path

import pytest

HERE = Path(__file__).resolve().parent
OFFICE = HERE.parent
REPO = OFFICE.parents[1]
EVAL = REPO / "tools" / "visual-eval"
FIX = EVAL / "fixtures"
sys.path.insert(0, str(OFFICE))
sys.path.insert(0, str(EVAL / "py"))

import chart_render  # noqa: E402
import doc_render  # noqa: E402
import measure  # noqa: E402
import numparse  # noqa: E402
import office  # noqa: E402
import pdf_render  # noqa: E402
import render_deck  # noqa: E402
import viz  # noqa: E402

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@pytest.fixture(autouse=True)
def _scratch(tmp_path, monkeypatch):
    monkeypatch.setenv("PI_OFFICE_GEN_SCRATCH", str(tmp_path / "scratch"))
    monkeypatch.setenv("PI_OFFICE_GEN_OFFLINE", "1")


def replies(name: str) -> dict:
    return json.loads((FIX / "replies" / f"{name}.json").read_text())


def deck_spec(r: dict) -> dict:
    """The spec office.make builds from a plan and its fills."""
    plan = r["plan"]
    slides = [{**f, "layout": p["layout"]} for p, f in zip(plan["slides"], r["fills"])]
    return {"theme": plan.get("theme", "ink"), "running_title": plan.get("running_title", ""), "slides": slides}


def captured(name: str) -> dict:
    return json.loads((FIX / "captured-4b" / name).read_text())


# ── D3 / D4 / D5: what a model writes, parsed ────────────────────────────────
@pytest.mark.parametrize("raw,value", [
    ("22M", 22e6), ("$38k", 38e3), ("1.2B", 1.2e9), ("$24,000", 24e3), ("3,100", 3100),
    ("3.1k", 3100), ("18%", 18), ("(12)", -12), ("−5", -5), ("1,200 GW", 1200), ("2.5B+", 2.5e9),
    (12, 12), ("$5m", 5e6),
])
def test_numbers_parse_as_written(raw, value):
    assert numparse.num(raw) == pytest.approx(value)


@pytest.mark.parametrize("raw", ["N/A", "2h 14m", "10–12", "Q3", "#1", ""])
def test_things_that_are_not_one_number_say_so(raw):
    assert numparse.parse(raw) is None


def test_string_booleans():
    assert numparse.truthy("false") is False and numparse.truthy("true") is True
    assert numparse.truthy("current") is False          # raw_fill5's "emphasis": "current"
    assert numparse.truthy("High", loose=True) is True   # a word in a matrix cell


def test_a_unit_only_display_keeps_its_number():
    assert numparse.shown("22M", "units") == "22M units"
    assert numparse.shown("1,200", "GW") == "1,200 GW"
    assert numparse.shown("3,100", "3,100 units") == "3,100 units"


def test_chart_render_reads_scaled_values():
    warn: list[str] = []
    spec = chart_render.normalise({"type": "bar", "items": [{"label": "a", "value": "22M"},
                                                            {"label": "b", "value": "N/A"}]}, warn)
    assert [p["value"] for p in spec["series"][0]["points"]] == [22e6, 0.0]
    assert any("'N/A'" in w for w in warn)


# ── the 4B-style pitch deck (research §2.2.1 a) ───────────────────────────────
@pytest.fixture(scope="module")
def deck_a(tmp_path_factory):
    out = tmp_path_factory.mktemp("deck") / "a.pptx"
    os.environ["PI_OFFICE_GEN_SCRATCH"] = str(out.parent / "scratch")
    spec = deck_spec(replies("pitch-deck-a"))
    drawn, warn = [], []
    render_deck.build(spec, out, drawn=drawn, warnings=warn)
    return out, spec, drawn, warn


def test_no_text_box_intersections_in_the_4b_deck(deck_a):
    out, spec, _, _ = deck_a
    m = measure.measure_pptx(out, spec)
    assert m["overlaps"] == [], m["overlaps"]           # D1 (title/subtitle), D2 ("3,10 / 0")
    assert m["offcanvas"] == []
    assert m["dropped"] == []


def _texts(slide):
    return [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]


def test_numerals_stay_on_one_line(deck_a):
    from pptx import Presentation
    out, *_ = deck_a
    sl = Presentation(str(out)).slides[3]
    nums = [s for s in _texts(sl) if s.text_frame.text in ("42", "3,100", "$38k", "18%")]
    assert len(nums) == 4
    for s in nums:
        assert len(s.text_frame.paragraphs) == 1 and s.text_frame.word_wrap is False


def test_market_bars_are_proportional_and_show_values(deck_a):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    out, *_ = deck_a
    sl = Presentation(str(out)).slides[4]
    labels = {s.text_frame.text for s in _texts(sl)}
    assert {"22M units", "3,100 units"} <= labels       # D4: the unit no longer replaces the number
    bars = [s for s in sl.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and not s.has_text_frame
            or (s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and not s.text_frame.text)]
    # track + fill per row; the fills are the ones in accent/primary, narrower than their track
    rows = sorted([b for b in bars if b.height < 400000 and b.width > 0], key=lambda b: (b.top, -b.width))
    tracks = {}
    for b in rows:
        tracks.setdefault(round(b.top / 50000), []).append(b.width)
    widths = [sorted(v) for v in tracks.values() if len(v) == 2]
    assert len(widths) == 2
    full, stub = sorted(widths, key=lambda v: v[0], reverse=True)
    assert full[0] == pytest.approx(full[1], rel=0.01)      # 22M fills its track (D3)
    assert stub[0] < stub[1] * 0.05                          # 3,100 of 22M is a stub


def test_emphasis_only_where_true(deck_a):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
    out, spec, *_ = deck_a
    sl = Presentation(str(out)).slides[2]
    t = render_deck.THEMES[spec["theme"]]
    nodes = [s for s in sl.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
             and s.auto_shape_type == MSO_SHAPE.OVAL]
    filled = [n for n in nodes if n.fill.fore_color.rgb == t.accent]
    assert len(nodes) == 3 and len(filled) == 1              # D5: "false" is false


# ── D6 / D19: truncation is reported, and the summary says what was drawn ────
def test_truncation_is_a_warning_and_the_summary_says_5_of_10(tmp_path):
    spec = captured("last_pptx_spec.json")                   # the REAL solar deck: 10 bars
    drawn, warn = [], []
    render_deck.build(spec, tmp_path / "solar.pptx", drawn=drawn, warnings=warn)
    assert any("5 of 10 shown" in w for w in warn)
    assert any("highlight_index 9" in w for w in warn)
    summary = office.summary_of("pptx", spec, drawn)
    assert "5 of 10 shown" in summary
    assert "1,200 GW" in summary                              # what is ON the bar, not "GW"


def test_flow_steps_past_five_are_reported(tmp_path):
    spec = deck_spec(replies("flow-diagram-a3"))
    warn: list[str] = []
    render_deck.build(spec, tmp_path / "f.pptx", warnings=warn)
    assert any("steps — 5 of 7 shown" in w for w in warn)


# ── D7: a failing layout leaves no half slide ─────────────────────────────────
def test_a_failing_layout_leaves_no_half_slide(tmp_path, monkeypatch):
    from pptx import Presentation

    def boom(sl, t, *a, **k):
        sl.shapes.add_textbox(0, 0, 100, 100).text_frame.text = "HALF-BUILT"
        raise RuntimeError("viz exploded")

    monkeypatch.setattr(viz, "bar_rows", boom)
    spec = deck_spec(replies("pitch-deck-a"))
    warn: list[str] = []
    out = tmp_path / "d.pptx"
    render_deck.build(spec, out, warnings=warn)
    prs = Presentation(str(out))
    assert len(prs.slides) == len(spec["slides"])            # no extra fallback slide
    text = " ".join(s.text_frame.text for s in _texts(prs.slides[4]))
    assert "HALF-BUILT" not in text and "22M" in text        # the fallback kept the content
    assert any("slide 5 (bars): could not be drawn" in w for w in warn)


# ── D8: a blank table header ─────────────────────────────────────────────────
def test_a_blank_table_header_does_not_crash(tmp_path):
    from pptx import Presentation
    spec = {"theme": "ink", "slides": [{"layout": "table", "title": "Market", "headers": ["", "Today", "Opportunity"],
                                        "rows": [["Units", "3,100", "22M"], ["Revenue", "$38k MRR"]]}]}
    warn: list[str] = []
    out = tmp_path / "t.pptx"
    render_deck.build(spec, out, warnings=warn)
    sl = Presentation(str(out)).slides[0]
    tables = [s for s in sl.shapes if s.has_table]
    assert len(tables) == 1 and not warn
    assert tables[0].table.cell(2, 0).text == "Revenue"


# ── the Word document (research §2.2.2) ───────────────────────────────────────
@pytest.fixture(scope="module")
def doc_a(tmp_path_factory):
    out = tmp_path_factory.mktemp("doc") / "a.docx"
    doc_render.build(replies("one-page-report-a")["doc"], out)
    return out


def _tables(path):
    from docx import Document
    return Document(str(path)).tables


def test_every_table_states_its_width_in_schema_order(doc_a):
    for t in _tables(doc_a):
        pr = t._tbl.tblPr
        tblw = pr.find(f"{W}tblW")
        assert tblw is not None and tblw.get(f"{W}type") == "dxa"
        assert pr.find(f"{W}tblLayout").get(f"{W}type") == "fixed"
        grid = sum(int(g.get(f"{W}w")) for g in t._tbl.find(f"{W}tblGrid"))
        assert abs(grid - int(tblw.get(f"{W}w"))) <= 2
        order = [c.tag.split("}")[1] for c in pr]
        seq = doc_render.TBLPR_SEQ
        assert order == sorted(order, key=seq.index), order   # D13: tblW no longer after tblLook


def test_stats_callout_and_table_span_the_text_width(doc_a):
    widths = [int(t._tbl.tblPr.find(f"{W}tblW").get(f"{W}w")) for t in _tables(doc_a)]
    assert widths and all(w == int(6.5 * 1440) for w in widths)


def test_stat_cells_are_padded_apart(doc_a):
    stats = next(t for t in _tables(doc_a) if len(t.columns) == 4)
    for cell in stats.rows[0].cells:
        mar = cell._tc.tcPr.find(f"{W}tcMar")
        assert mar is not None and int(mar.find(f"{W}right").get(f"{W}w")) >= 0.2 * 1440


def test_heading_rules_are_at_most_6pt(doc_a):
    rules = measure.measure_docx(doc_a)["rules"]
    assert rules and all(r["height_pt"] <= 6 for r in rules)            # D14 (was a ~13 pt tick)
    assert all(r["width_in"] == pytest.approx(0.65, abs=0.02) for r in rules)


def test_the_cover_band_is_padded(doc_a):
    cover = _tables(doc_a)[0]
    mar = cover.cell(0, 0)._tc.tcPr.find(f"{W}tcMar")
    assert int(mar.find(f"{W}left").get(f"{W}w")) >= 0.3 * 1440          # D15


def test_the_q3_b_spec_fits_one_page(tmp_path):
    out = tmp_path / "b.docx"
    doc_render.build(replies("one-page-report-b")["doc"], out)
    assert measure.measure_docx(out)["pages_est"] == 1


def test_a_failing_block_is_rolled_back(tmp_path, monkeypatch):
    from docx import Document

    def boom(doc, t, b):
        doc.add_paragraph("HALF-WRITTEN")
        raise RuntimeError("block exploded")

    monkeypatch.setitem(doc_render.BLOCKS, "quote", boom)
    warn: list[str] = []
    out = tmp_path / "r.docx"
    doc_render.build(replies("one-page-report-a")["doc"], out, warnings=warn)
    assert "HALF-WRITTEN" not in "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert any("(quote): could not be drawn" in w for w in warn)


def test_docx_truncation_is_reported(tmp_path):
    spec = {"blocks": [{"type": "bullets", "items": [f"item {i}" for i in range(11)]}]}
    drawn, warn = [], []
    doc_render.build(spec, tmp_path / "x.docx", drawn=drawn, warnings=warn)
    assert any("items — 8 of 11 shown" in w for w in warn)
    assert "8 of 11 shown" in office.summary_of("docx", spec, drawn)


# ── the PDF ──────────────────────────────────────────────────────────────────
def test_pdf_cover_mid_document_does_not_paint_over_content(tmp_path):
    import pypdf
    out = tmp_path / "u.pdf"
    pdf_render.build(captured("last_pdf_spec.json"), out)
    pages = [p.extract_text() or "" for p in pypdf.PdfReader(str(out)).pages]
    outlook = [i for i, t in enumerate(pages) if "Strategic Outlook" in t]
    assert outlook, pages
    assert "resilience" not in pages[outlook[0]]            # the quote is not under the cover
    assert any("resilience" in t for t in pages)
    assert pages[-1].strip() and "Strategic Outlook" in pages[-1]   # no blank page after it


def test_pdf_bars_and_table_cells(tmp_path):
    import pypdf
    spec = {"blocks": [
        {"type": "bars", "items": [{"label": "Market", "value": "22M", "display": "units"},
                                   {"label": "Today", "value": "3,100", "display": "units"}]},
        {"type": "table", "headers": ["Region", "A very long header that used to be cut at twenty-six"],
         "rows": [["North", "a cell with more than thirty characters in it, kept whole"]]},
    ]}
    out = tmp_path / "b.pdf"
    pdf_render.build(spec, out)
    text = " ".join((p.extract_text() or "") for p in pypdf.PdfReader(str(out)).pages).replace("\n", " ")
    assert "22M units" in text and "3,100 units" in text
    assert "kept whole" in text and "twenty-six" in text


# ── office.py end to end, against the replay stub ────────────────────────────
def _free_port() -> int:
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    p = s.getsockname()[1]
    s.close()
    return p


def test_office_make_reports_what_it_drew(tmp_path):
    spec = captured("last_pptx_spec.json")
    rep = {"plan": {"theme": spec["theme"], "running_title": spec["running_title"],
                    "slides": [{"layout": s["layout"], "title": s.get("title", "")} for s in spec["slides"]]},
           "fills": spec["slides"]}
    rfile = tmp_path / "replies.json"
    rfile.write_text(json.dumps(rep))
    port = _free_port()
    srv = subprocess.Popen([sys.executable, str(EVAL / "py" / "replay_server.py"), str(rfile), str(port)])
    try:
        for _ in range(100):
            try:
                socket.create_connection(("127.0.0.1", port), timeout=0.2).close()
                break
            except OSError:
                time.sleep(0.05)
        env = {**os.environ, "PI_OFFICE_GEN_SERVER": f"http://127.0.0.1:{port}",
               "PI_OFFICE_GEN_SCRATCH": str(tmp_path / "s"), "no_proxy": "*", "NO_PROXY": "*"}
        brief = json.loads((FIX / "captured-4b" / "briefs.json").read_text())["solar-deck"]["brief"]
        r = subprocess.run([sys.executable, str(OFFICE / "office.py"), "make", "pptx", "--brief", brief,
                            "--out", str(tmp_path / "solar.pptx")], env=env, capture_output=True, text=True,
                           timeout=120)
    finally:
        srv.terminate()
        srv.wait(timeout=5)
    out = json.loads(r.stdout.strip().splitlines()[-1])
    assert out["ok"], out
    assert any("5 of 10 shown" in w for w in out["warnings"])
    assert "5 of 10 shown" in out["summary"]
    assert not r.stdout.strip().splitlines()[:-1]            # only the JSON on stdout
