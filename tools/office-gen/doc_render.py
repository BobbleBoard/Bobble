#!/usr/bin/env python3
"""
Spec (JSON) -> .docx, with the design owned here and the words owned by the model.

Same bet as the deck renderer: the model picks a block type from a short menu and
writes the prose; this file owns typography, colour, spacing and the page.

python-docx has no styling conveniences worth the name — shaded paragraphs,
borders and column widths all mean writing raw OOXML. That is why most generated
Word documents look like Times New Roman with bullets: the library makes the
plain thing easy and everything else fiddly. The `_shade`/`_border` helpers below
are the fiddly part, done once.

VQ-01 (2026-09-23), what the file looked like when opened (QuickLook; Word lays
tables out the same way):
  - every table now states its width (tblW), a fixed layout and its column
    grid. Setting only CELL widths does nothing — Word sizes a table from
    tblW/tblGrid — so the stats row collapsed to "12,4802h 14m91%640" and the
    callout sat at half width (D13);
  - a heading rule is a 3 pt paragraph border, not a 13 pt tick: it was a
    one-cell table whose paragraph MARK kept the 11 pt Normal size (D14), and
    a table's width is not honoured by every viewer (below);
  - the cover band has inner padding (D15);
  - every property is written where the OOXML schema's sequence puts it. The
    old tblW and tblBorders were APPENDED after tblLook, out of order, which a
    strict reader skips — part of why the width never took;
  - cells carry explicit padding, because a viewer that sizes a table by its
    content (QuickLook does, whatever tblW says — MEASURED) would otherwise
    set "12,480" against "2h 14m" and "Billing questions" against "31%";
  - lists longer than a block holds are cut AND reported, a failing block is
    rolled back and reported, and `build(drawn=…)` says what was drawn.

WF-06 (2026-09-24), the render path's document — `office.py render docx`, a
spec built by code from pages it fetched (workflows.md §4.5), never by the make
path's model. Given `cite=` (citations.py) the document can:
  - carry a `sources` block: the spec's sources, numbered, each a paragraph
    with a bookmark and its URL as a real external link;
  - set every `[S3]` marker as a superscript "3" that is an internal link to
    source 3 — so a reader clicks from the claim to where it came from;
  - carry a `toc` block: the headings, each an internal link to its bookmark.
Without `cite` (every make) none of it exists and nothing here runs: a small
model's spec does not get to add a bibliography (VQ-08).
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent))
import numparse  # noqa: E402
import palette as pal  # noqa: E402
import textfit  # noqa: E402

SANS = "Helvetica Neue"
SERIF = "Georgia"
TEXT_W = 6.5          # inches between the margins (8.5 in Letter, 1 in each side)


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#").upper())


def _s(v) -> str:
    return numparse.as_text(v)


# ── what was drawn ───────────────────────────────────────────────────────────
class BlockReport:
    def __init__(self, index: int, kind: str):
        self.index, self.kind = index, kind
        self.head = ""
        self.items: list[str] = []
        self.cuts: list[tuple[str, int, int]] = []
        self.notes: list[str] = []

    def as_dict(self) -> dict:
        return {"block": self.index, "type": self.kind, "head": self.head, "items": self.items,
                "cuts": [{"field": f, "given": g, "shown": s} for f, g, s in self.cuts],
                "notes": self.notes}


_R: BlockReport | None = None
# The render path only (WF-06): the spec's sources, and the (bookmark, title)
# of each block a contents list links to. None / empty on every make.
_C = None
_T = None             # the palette, for the links' colour
_TOC: dict[int, tuple[str, str]] = {}


def _cap(seq, n: int, field: str) -> list:
    items = list(seq) if isinstance(seq, (list, tuple)) else ([] if seq is None else [seq])
    if len(items) > n and _R is not None:
        _R.cuts.append((field, len(items), n))
    return items[:n]


# ── OOXML helpers ─────────────────────────────────────────────────────────────
# The child ORDER the schema requires in each properties element. A reader is
# entitled to skip an element that is out of sequence, and some do.
TBLPR_SEQ = ("tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
             "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
             "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription",
             "tblPrChange")
TCPR_SEQ = ("cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
            "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns",
            "cellDel", "cellMerge", "tcPrChange")
PPR_SEQ = ("pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
           "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku",
           "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi",
           "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
           "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap",
           "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange")
RPR_SEQ = ("rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike",
           "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden",
           "color", "spacing", "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect",
           "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout",
           "specVanish", "oMath")


def _put(parent, el, seq):
    """Insert `el` where the schema sequence puts it, replacing any of its kind."""
    name = el.tag.split("}")[1]
    for old in parent.findall(qn(f"w:{name}")):
        parent.remove(old)
    later = set(seq[seq.index(name) + 1:])
    for i, child in enumerate(parent):
        if child.tag.split("}")[1] in later:
            parent.insert(i, el)
            return el
    parent.append(el)
    return el


def _shd(hex_colour: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_colour.lstrip("#").upper())
    return shd


def _p_shade(p, hex_colour: str):
    _put(p._p.get_or_add_pPr(), _shd(hex_colour), PPR_SEQ)


def _cell_shade(cell, hex_colour: str):
    """Cell background. No python-docx API exists; this is the OOXML."""
    _put(cell._tc.get_or_add_tcPr(), _shd(hex_colour), TCPR_SEQ)


def _no_borders(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "none")
        e.set(qn("w:sz"), "0")
        borders.append(e)
    _put(table._tbl.tblPr, borders, TBLPR_SEQ)


def _tbl_width(table, inches: float, cols: list[float] | None = None):
    """Pin a table's total width, its column grid and every cell's width, with
    a FIXED layout. Setting only the CELL width does nothing — Word lays the
    table out from tblW/tblGrid, so a table without them collapses to its
    content (the stats row, D13) or runs full-bleed (the rule under a heading)."""
    ncols = len(table.columns)
    cols = cols or [inches / max(1, ncols)] * ncols
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(int(round(sum(cols) * 1440))))
    w.set(qn("w:type"), "dxa")
    _put(table._tbl.tblPr, w, TBLPR_SEQ)
    table.autofit = False          # python-docx writes tblLayout=fixed, in sequence
    grid = table._tbl.find(qn("w:tblGrid"))
    for gc, cw in zip(grid.findall(qn("w:gridCol")), cols):
        gc.set(qn("w:w"), str(int(round(cw * 1440))))
    for row in table.rows:
        for cell, cw in zip(row.cells, cols):
            cell.width = Inches(cw)


def _cell_margins(cell, *, left=None, right=None, top=None, bottom=None):
    """Inner padding of one cell, in inches (tcMar) — merged with any set before."""
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn("w:tcMar"))
    have = {}
    if old is not None:
        for e in old:
            have[e.tag.split("}")[1]] = e.get(qn("w:w"))
    for side, v in (("left", left), ("right", right), ("top", top), ("bottom", bottom)):
        if v is not None:
            have[side] = str(int(round(v * 1440)))
    mar = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        if side in have:
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:w"), have[side])
            e.set(qn("w:type"), "dxa")
            mar.append(e)
    _put(tcPr, mar, TCPR_SEQ)


def _mark_size(p, pt: float):
    """The paragraph MARK's size. An empty paragraph is one line of it, so a
    3 pt rule whose mark stays 11 pt is an 11 pt-tall rule."""
    rPr = OxmlElement("w:rPr")
    for tag in ("w:sz", "w:szCs"):
        e = OxmlElement(tag)
        e.set(qn("w:val"), str(max(2, int(pt * 2))))
        rPr.append(e)
    _put(p._p.get_or_add_pPr(), rPr, PPR_SEQ)


def _rule(doc, colour: str, *, width_pct=18, space_before=2, space_after=10):
    """The accent rule under a heading: an empty paragraph's BOTTOM BORDER,
    3 pt thick, as wide as a right indent leaves it. It used to be a one-cell
    shaded table, which rendered as a 13 pt tick: the cell's paragraph mark set
    its height, and a viewer that sizes tables by content drew it one
    character wide (D14). A paragraph border has neither problem."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = Pt(1)                       # an EXACT 1 pt line
    pf.right_indent = Inches(TEXT_W * (1 - width_pct / 100))
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "24")                  # eighths of a point: 3 pt
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), colour.lstrip("#").upper())
    bdr.append(bottom)
    _put(p._p.get_or_add_pPr(), bdr, PPR_SEQ)
    _mark_size(p, 1)
    return p


def _spacer(doc, after_pt: float):
    """The gap after a table: `after_pt` of space and nothing else. It was an
    empty 11 pt Normal paragraph PLUS the space — about 13 pt of blank line
    more than asked for, after every table, which is how a one-page report
    spilled onto a second page."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(after_pt)
    pf.line_spacing = Pt(1)
    _mark_size(p, 1)
    return p


def _run(p, s, *, size=11, colour=None, bold=False, italic=False, font=SANS,
         spacing=None):
    r = p.add_run(_s(s))
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    if colour is not None:
        r.font.color.rgb = _rgb(colour)
    if spacing is not None:
        el = OxmlElement("w:spacing")
        el.set(qn("w:val"), str(int(spacing * 20)))
        _put(r._r.get_or_add_rPr(), el, RPR_SEQ)
    return r


def _text(p, s, *, size=11, colour=None, bold=False, italic=False, font=SANS,
          spacing=None):
    """One run of text — or, on the render path, the text around its citation
    markers with each marker set as linked superscripts."""
    fmt = {"size": size, "colour": colour, "bold": bold, "italic": italic, "font": font,
           "spacing": spacing}
    if _C is None:
        return _run(p, s, **fmt)
    first = None
    for seg in _C.split(_s(s)):
        r = _run(p, seg, **fmt) if isinstance(seg, str) else _cite(p, seg, size=size, font=font)
        first = first if first is not None else r
    return first


# ── links and anchors (render path) ──────────────────────────────────────────
# Bookmark names start with "_": Word, GenOffice and mammoth all treat those as
# document-internal anchors (Word's own _Ref/_Toc), not as bookmarks a person
# made — so they never clutter a bookmark list, and a link to one still works.
def _source_anchor(n: int) -> str:
    return f"_RefSource{n}"


_BOOKMARK_ID = [0]


def _bookmark(p, name: str) -> None:
    """Mark the whole paragraph as `name`, the target of an internal link."""
    _BOOKMARK_ID[0] += 1
    bid = str(_BOOKMARK_ID[0])
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bid)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bid)
    ppr = p._p.pPr
    if ppr is not None:
        ppr.addnext(start)
    else:
        p._p.insert(0, start)
    p._p.append(end)


def _link(p, text: str, *, url: str | None = None, anchor: str | None = None, size=11,
          colour=None, bold=False, underline=False, superscript=False, font=SANS):
    """A run inside a w:hyperlink: an external URL (a relationship, TargetMode
    External) or an internal anchor. Styled directly — the built-in Hyperlink
    style would bring Word's followed-link purple with it."""
    r = _run(p, text, size=size, colour=colour, bold=bold, font=font)
    if underline:
        r.font.underline = True
    if superscript:
        r.font.superscript = True
    h = OxmlElement("w:hyperlink")
    if url:
        h.set(qn("r:id"), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    else:
        h.set(qn("w:anchor"), anchor or "")
    h.set(qn("w:history"), "1")
    r._r.addprevious(h)
    h.append(r._r)
    return r


def _cite(p, nums: list[int], *, size=11, font=SANS):
    """Superscript source numbers, each a link to its entry: "rose 14%¹˒³"."""
    colour = _T.primary if _T is not None else None
    first = None
    for i, n in enumerate(nums):
        if i:
            sep = _run(p, ",", size=size, colour=colour, font=font)
            sep.font.superscript = True
        r = _link(p, str(n), anchor=_source_anchor(n), size=size, colour=colour, superscript=True,
                  font=font)
        first = first if first is not None else r
    return first


def _para(doc, *, before=0, after=8, align=None, line=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if line is not None:
        pf.line_spacing = line
    return p


# ── blocks ────────────────────────────────────────────────────────────────────
def B_cover(doc, t, b):
    """A full-bleed colour band with the title reversed out of it — padded, so
    the title no longer runs to the band's edge (D15)."""
    tbl = doc.add_table(rows=1, cols=1)
    _no_borders(tbl)
    _tbl_width(tbl, TEXT_W)
    cell = tbl.cell(0, 0)
    _cell_shade(cell, t.deep)
    _cell_margins(cell, left=0.4, right=0.4)
    fg = t.on(t.deep)
    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_before = Pt(54)
    p0.paragraph_format.space_after = Pt(6)
    if b.get("eyebrow"):
        _text(p0, _s(b["eyebrow"]).upper(), size=9, colour=t.accent, bold=True, spacing=1.6)
    p1 = cell.add_paragraph()
    p1.paragraph_format.space_after = Pt(10)
    _text(p1, b.get("title", ""), size=30, colour=fg, bold=True)
    if b.get("subtitle"):
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(60)
        _text(p2, b["subtitle"], size=12.5, colour=t.support)
    else:
        p1.paragraph_format.space_after = Pt(60)
    _spacer(doc, 18)
    if b.get("byline"):
        p = _para(doc, after=16)
        _text(p, b["byline"], size=9.5, colour=t.mute, spacing=0.8)
    if _R is not None:
        _R.head = _s(b.get("title"))


def B_heading(doc, t, b):
    p = _para(doc, before=16, after=2)
    _text(p, b.get("title", ""), size=17, colour=t.primary, bold=True)
    if _R is not None and _R.index in _TOC:
        _bookmark(p, _TOC[_R.index][0])    # a contents entry links here
    _rule(doc, t.accent, width_pct=10, space_after=4)
    if b.get("standfirst"):
        p2 = _para(doc, after=10, line=1.35)
        _text(p2, b["standfirst"], size=11.5, colour=t.mute, italic=True)
    if _R is not None:
        _R.head = _s(b.get("title"))
        _R.items = [_s(b.get("standfirst"))] if b.get("standfirst") else []


def B_body(doc, t, b):
    paras = _cap(b.get("paragraphs") or [], 6, "paragraphs")
    for para in paras:
        p = _para(doc, after=9, line=1.4)
        _text(p, para, size=11, colour=t.ink)
    if _R is not None:
        _R.items = [_s(x) for x in paras]


def B_callout(doc, t, b):
    """A tinted panel with an accent spine — the highest-contrast way to say
    'this bit matters' without shouting."""
    tbl = doc.add_table(rows=1, cols=2)
    _no_borders(tbl)
    _tbl_width(tbl, TEXT_W, [0.06, TEXT_W - 0.06])
    spine, body = tbl.cell(0, 0), tbl.cell(0, 1)
    _cell_shade(spine, t.accent)
    _cell_shade(body, t.faint)
    _cell_margins(spine, left=0, right=0)
    p = body.paragraphs[0]
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.16)
    if b.get("label"):
        _text(p, _s(b["label"]).upper(), size=8.5, colour=t.accent, bold=True, spacing=1.4)
        p = body.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.right_indent = Inches(0.16)
    _text(p, b.get("text", ""), size=11.5, colour=t.ink, bold=True)
    _spacer(doc, 12)
    if _R is not None:
        _R.head = _s(b.get("label"))
        _R.items = [_s(b.get("text"))]


def B_stats(doc, t, b):
    """A row of oversized numerals. The only reliable way to make a Word page
    stop looking like a Word page."""
    items = _cap(b.get("stats") or [], 4, "stats")
    if not items:
        return
    n = len(items)
    col = TEXT_W / n
    values = [_s(it.get("value", "") if isinstance(it, dict) else it) for it in items]
    labels = [_s(it.get("label", "") if isinstance(it, dict) else "") for it in items]
    # One size for the row: the largest at which every value fits its column.
    size = 16
    for cand in (26, 24, 22, 20, 18, 16):
        size = cand
        if all(textfit.width_pt(v, cand, bold=True) <= (col - 0.3) * 72 for v in values):
            break
    tbl = doc.add_table(rows=2, cols=n)
    _no_borders(tbl)
    _tbl_width(tbl, TEXT_W)
    for i, (value, label) in enumerate(zip(values, labels)):
        top, bot = tbl.cell(0, i), tbl.cell(1, i)
        for c in (top, bot):
            _cell_margins(c, left=0, right=0.22)
        pv = top.paragraphs[0]
        pv.paragraph_format.space_after = Pt(0)
        _text(pv, value, size=size, colour=t.primary, bold=True)
        pl = bot.paragraphs[0]
        pl.paragraph_format.space_after = Pt(12)
        _text(pl, label, size=9, colour=t.mute)
    _spacer(doc, 10)
    if _R is not None:
        _R.items = [f"{v} {lab}".strip() for v, lab in zip(values, labels)]


def B_table(doc, t, b):
    headers = [_s(h) for h in (b.get("headers") or [])]
    rows = [r if isinstance(r, list) else [r] for r in _cap(b.get("rows") or [], 12, "rows")]
    ncols = max([len(headers)] + [len(r) for r in rows] + [0])
    if ncols == 0:
        return
    headers += [""] * (ncols - len(headers))
    # Columns as wide as their content asks, within the text width.
    need = []
    for c in range(ncols):
        texts = [headers[c]] + [_s(r[c]) for r in rows if c < len(r)]
        need.append(max(0.8, min(3.2, max(textfit.width_pt(x, 10, bold=(c == 0)) for x in texts) / 72 + 0.25)))
    scale = TEXT_W / sum(need)
    widths = [w * scale for w in need]
    tbl = doc.add_table(rows=len(rows) + 1, cols=ncols)
    _no_borders(tbl)
    _tbl_width(tbl, TEXT_W, widths)
    for row in tbl.rows:
        for cell in row.cells:
            _cell_margins(cell, left=0.1, right=0.1)
    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        _cell_shade(cell, t.deep)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(5)
        _text(p, h, size=9.5, colour=t.on(t.deep), bold=True)
    for r, row in enumerate(rows, start=1):
        for c in range(ncols):
            cell = tbl.cell(r, c)
            if r % 2 == 0:
                _cell_shade(cell, t.faint)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            val = row[c] if c < len(row) else ""
            _text(p, val, size=10, colour=t.ink, bold=(c == 0))
    _spacer(doc, 14)
    if _R is not None:
        _R.head = " · ".join(h for h in headers if h)
        _R.items = [_s(r[0]) if r else "" for r in rows]


def B_quote(doc, t, b):
    p = _para(doc, before=10, after=4, line=1.3)
    p.paragraph_format.left_indent = Inches(0.35)
    _text(p, "“" + _s(b.get("quote", "")) + "”", size=15, colour=t.primary,
          italic=True, font=SERIF)
    if b.get("attribution"):
        p2 = _para(doc, after=14)
        p2.paragraph_format.left_indent = Inches(0.35)
        _text(p2, "— " + _s(b["attribution"]), size=9.5, colour=t.mute)
    if _R is not None:
        _R.head = _s(b.get("quote"))


def B_bullets(doc, t, b):
    items = _cap(b.get("items") or [], 8, "items")
    for item in items:
        p = _para(doc, after=6, line=1.35)
        p.paragraph_format.left_indent = Inches(0.22)
        _text(p, "▪  ", size=10, colour=t.accent, bold=True)
        _text(p, item, size=11, colour=t.ink)
    if _R is not None:
        _R.items = [_s(x) for x in items]


def B_pagebreak(doc, t, b):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ── render path only (WF-06) ─────────────────────────────────────────────────
NUM_W = 0.34          # inches: the number column of a sources entry


def B_sources(doc, t, b):
    """The spec's sources, in its order and numbered as the citations number
    them: title and site, then the URL as a live external link. Each entry is
    a bookmark, so a superscript in the text lands on it."""
    B_heading(doc, t, {"title": _s(b.get("title")) or "Sources"})
    items = []
    for src in _C.sources:
        p = _para(doc, after=7, line=1.2)
        pf = p.paragraph_format
        pf.left_indent = Inches(NUM_W)
        pf.first_line_indent = Inches(-NUM_W)
        pf.tab_stops.add_tab_stop(Inches(NUM_W))
        _bookmark(p, _source_anchor(src.n))
        _run(p, f"{src.n}\t", size=10, colour=t.accent, bold=True)
        _run(p, src.title, size=10.5, colour=t.ink, bold=True)
        tail = ", ".join(x for x in (src.site, src.date) if x)
        if tail:
            _run(p, f" — {tail}", size=10, colour=t.mute)
        if src.url or src.accessed:
            p.add_run().add_break()
        if src.url:
            _link(p, src.url, url=src.url, size=9.5, colour=t.primary, underline=True)
        if src.accessed:
            _run(p, f"{'  ·  ' if src.url else ''}accessed {src.accessed}", size=9, colour=t.mute)
        items.append(f"{src.n}. {src.label()}" + (f" {src.url}" if src.url else ""))
    if _R is not None:
        _R.items = items


def B_toc(doc, t, b):
    """Contents: every heading (and the sources) as an internal link to its
    bookmark. No page numbers — a viewer that does not lay the document out
    cannot know them, and a wrong page number is worse than none."""
    B_heading(doc, t, {"title": _s(b.get("title")) or "Contents"})
    for _i, (anchor, title) in sorted(_TOC.items()):
        p = _para(doc, after=4, line=1.2)
        _link(p, title, anchor=anchor, size=11, colour=t.primary)
    if _R is not None:
        _R.items = [title for _a, title in (_TOC[i] for i in sorted(_TOC))]


BLOCKS = {
    "cover": B_cover, "heading": B_heading, "body": B_body, "callout": B_callout,
    "stats": B_stats, "table": B_table, "quote": B_quote, "bullets": B_bullets,
    "pagebreak": B_pagebreak,
}
RENDER_BLOCKS = {"sources": B_sources, "toc": B_toc}


def _render_plan(blocks: list, cite) -> list:
    """The render path's block list: the sources appended when the spec has
    sources but placed no block for them (a citation must never point at
    nothing), and a bookmark for every heading when there is a contents list."""
    global _TOC
    kinds = [_s(b.get("type")) if isinstance(b, dict) else "body" for b in blocks]
    if cite.sources and "sources" not in kinds:
        blocks = blocks + [{"type": "sources"}]
        kinds.append("sources")
    _TOC = {}
    if "toc" in kinds:
        for i, (b, kind) in enumerate(zip(blocks, kinds), 1):
            if kind == "heading":
                _TOC[i] = (f"_TocSection{i}", _s(b.get("title")))
            elif kind == "sources":
                _TOC[i] = (f"_TocSection{i}", _s(b.get("title")) or "Sources")
    return blocks


def build(spec: dict, out: Path, drawn: list | None = None, warnings: list | None = None, *,
          cite=None) -> Path:
    """Render the document. `cite` (citations.Citations) is the render path's:
    given, the spec's sources, citations and contents are drawn; make never
    passes it."""
    global _R, _C, _T, _TOC
    t = pal.from_spec(spec)
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.9)
        s.left_margin = s.right_margin = Inches(1.0)
    normal = doc.styles["Normal"]
    normal.font.name = SANS
    normal.font.size = Pt(11)

    blocks = list(spec.get("blocks", []))
    registry = BLOCKS
    if cite is not None:
        import citations
        _C, _T = cite, t
        _BOOKMARK_ID[0] = 0
        blocks = _render_plan(blocks, cite)
        registry = {**BLOCKS, **RENDER_BLOCKS}
        citations.theme_links(doc.part, t.primary)
    try:
        _draw(doc, t, blocks, registry, drawn, warnings)
    finally:
        _R, _C, _T, _TOC = None, None, None, {}
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def _draw(doc, t, blocks: list, registry: dict, drawn: list | None, warnings: list | None) -> None:
    global _R
    body = doc.element.body
    for i, b in enumerate(blocks, 1):
        if not isinstance(b, dict):
            b = {"type": "body", "paragraphs": [_s(b)]}
        kind = _s(b.get("type")) or "body"
        _R = BlockReport(i, kind)
        fn = registry.get(kind)
        if fn is None:
            _R.notes.append(f"block type '{kind}' does not exist; set as body text")
            fn = B_body
            b = {**b, "paragraphs": b.get("paragraphs") or [b.get("text", "")]}
        before = len(body)
        try:
            fn(doc, t, b)
        except Exception as err:  # noqa: BLE001
            # Roll the half-written block back out of the body (the sectPr is
            # always last, so everything between is this block's).
            added = list(body)[before - 1:-1] if before else []
            for el in added:
                body.remove(el)
            print(f"  ! block {i} ({kind}) failed: {err}", file=sys.stderr)
            _R.notes.append(f"could not be drawn ({type(err).__name__}: {err}); left out")
        if drawn is not None:
            drawn.append(_R.as_dict())
        if warnings is not None:
            for field, given, shown in _R.cuts:
                warnings.append(f"block {i} ({kind}): {field} — {shown} of {given} shown; "
                                f"the block holds {shown}, put the rest in another block")
            for note in _R.notes:
                warnings.append(f"block {i} ({kind}): {note}")


if __name__ == "__main__":
    spec = json.loads(Path(sys.argv[1]).read_text())
    warn: list[str] = []
    p = build(spec, Path(sys.argv[2]), warnings=warn)
    print(f"wrote {p} ({len(spec.get('blocks', []))} blocks)")
    for w in warn:
        print(f"  ! {w}")
