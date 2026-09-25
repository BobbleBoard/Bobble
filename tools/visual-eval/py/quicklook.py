"""
Render office files to PNGs with QuickLook, headless: `qlmanage -t` is a
command-line thumbnailer and never shows a window or takes focus.

QuickLook draws only the FIRST slide of a deck and the first page of a PDF, so
each slide / page is cut into its own one-slide deck / one-page PDF first. A
.docx gets its first page (what "does it fit on one page" is about). Every cut
file of every artifact goes to ONE qlmanage call — it takes many paths, and
the whole eval's office renders then cost about a second.

    python quicklook.py manifest.json     # [{"src", "out_dir", "prefix"}] -> JSON list of PNGs per item

QuickLook is macOS-only and belongs to the eval harness, never to the app's
runtime (visual-quality.md §4.8, track 4).
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

SIZE = 1600


def split_pptx(src: Path, work: Path, prefix: str) -> list[Path]:
    from pptx import Presentation
    n = len(Presentation(str(src)).slides)
    outs = []
    for i in range(n):
        prs = Presentation(str(src))
        lst = prs.slides._sldIdLst
        for j, sid in enumerate(list(lst)):
            if j != i:
                lst.remove(sid)
        p = work / f"{prefix}__slide{i + 1:02d}.pptx"
        prs.save(str(p))
        outs.append(p)
    return outs


def split_pdf(src: Path, work: Path, prefix: str) -> list[Path]:
    import pypdf
    reader = pypdf.PdfReader(str(src))
    outs = []
    for i, page in enumerate(reader.pages):
        w = pypdf.PdfWriter()
        w.add_page(page)
        p = work / f"{prefix}__page{i + 1:02d}.pdf"
        with open(p, "wb") as fh:
            w.write(fh)
        outs.append(p)
    return outs


def render(items: list[dict]) -> list[list[str]]:
    work = Path(tempfile.mkdtemp(prefix="vq-ql-"))
    try:
        plan: list[list[tuple[Path, Path]]] = []
        cuts: list[Path] = []
        for it in items:
            src, out_dir, prefix = Path(it["src"]), Path(it["out_dir"]), it["prefix"]
            out_dir.mkdir(parents=True, exist_ok=True)
            ext = src.suffix.lower()
            if ext == ".pptx":
                parts = split_pptx(src, work, prefix)
                names = [out_dir / f"{prefix}-slide{k:02d}.png" for k in range(1, len(parts) + 1)]
            elif ext == ".pdf":
                parts = split_pdf(src, work, prefix)
                names = [out_dir / f"{prefix}-page{k:02d}.png" for k in range(1, len(parts) + 1)]
            else:
                one = work / f"{prefix}__page01{ext}"
                shutil.copy(src, one)
                parts = [one]
                names = [out_dir / f"{prefix}-page01.png"]
            plan.append(list(zip(parts, names)))
            cuts.extend(parts)
        if cuts:
            subprocess.run(["qlmanage", "-t", "-s", str(SIZE), "-o", str(work), *map(str, cuts)],
                           capture_output=True, timeout=300)
        out: list[list[str]] = []
        for pairs in plan:
            got = []
            for part, name in pairs:
                png = work / (part.name + ".png")
                if png.exists() and png.stat().st_size > 2000:
                    shutil.move(str(png), name)
                    got.append(str(name))
            out.append(got)
        return out
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    items = json.loads(Path(sys.argv[1]).read_text())
    print(json.dumps(render(items)))
