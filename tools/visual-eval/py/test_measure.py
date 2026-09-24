"""
The eval's ruler, tested on files built for the purpose: a measurement that
cannot see a known collision (or sees one that is not there) would make every
report it writes worthless.

    tools/office-gen/.venv/bin/python -m pytest tools/visual-eval/py -q
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import measure  # noqa: E402
from palette_check import categorical, contrast, delta_e, parse_colour, simulate  # noqa: E402

pptx = pytest.importorskip("pptx")
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402


def _deck(tmp_path, boxes, name="t.pptx"):
    """boxes: [(x_in, y_in, w_in, h_in, text, size_pt, wrap)] on one 13.333 x 7.5 slide."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    for x, y, w, h, text, size, wrap in boxes:
        tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = wrap
        for m in ("left", "right", "top", "bottom"):
            setattr(tf, f"margin_{m}", 0)
        r = tf.paragraphs[0].add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.name = "Helvetica Neue"
        r.font.color.rgb = RGBColor(0x14, 0x14, 0x19)
    p = tmp_path / name
    prs.save(str(p))
    return p


def test_a_title_that_wraps_into_the_subtitle_is_an_overlap(tmp_path):
    # hero.py's D1: 58 pt in a 7.2 in box wraps to four lines; the subtitle sits at a fixed 5.05 in.
    p = _deck(tmp_path, [
        (0.9, 2.15, 7.2, 3.0, "Tidewell: Smart Water-Leak Sensors for Apartment Buildings", 58, True),
        (0.9, 5.05, 6.1, 1.0, "Stopping water damage before it starts", 16.5, True),
    ])
    m = measure.measure_pptx(p)
    assert len(m["overlaps"]) == 1


def test_stacked_boxes_with_room_do_not_overlap(tmp_path):
    p = _deck(tmp_path, [
        (0.9, 1.0, 11, 1.0, "A short title", 34, True),
        (0.9, 2.0, 11, 1.0, "A subtitle well below it", 17, True),
    ])
    assert measure.measure_pptx(p)["overlaps"] == []


def test_a_numeral_wider_than_its_column_breaks_over_its_label(tmp_path):
    # render_deck.L_stats's D2: "3,100" at a fixed 76 pt in a 2.5 in column.
    p = _deck(tmp_path, [
        (3.5, 4.0, 2.5, 1.5, "3,100", 76, True),
        (3.5, 5.6, 2.3, 1.0, "Units", 12.5, True),
    ])
    assert len(measure.measure_pptx(p)["overlaps"]) == 1


def test_small_text_share_and_fonts(tmp_path):
    p = _deck(tmp_path, [(1, 1, 5, 1, "abcd", 12, True), (1, 3, 5, 1, "efghijkl", 20, True)])
    m = measure.measure_pptx(p)
    assert m["text"]["small_pct"] == pytest.approx(33.3, abs=0.1)
    assert m["fonts"] == ["Helvetica Neue"]


def test_content_on_one_third_leaves_an_empty_region(tmp_path):
    p = _deck(tmp_path, [(0.5, 0.5, 3.5, 6.5, "word " * 120, 14, True)])
    m = measure.measure_pptx(p)
    assert m["empty_regions"] and m["empty_regions"][0]["pct"] > 50


def test_dropped_items_are_counted_against_the_spec(tmp_path):
    p = _deck(tmp_path, [(1, 1, 8, 1, "2020", 14, True), (1, 2, 8, 1, "2021", 14, True)])
    spec = {"slides": [{"layout": "bars", "items": [{"label": "2020"}, {"label": "2021"}, {"label": "2022"}]}]}
    d = measure.measure_pptx(p, spec)["dropped"]
    assert d == [{"slide": 1, "layout": "bars", "field": "items", "given": 3, "drawn": 2, "missing": ["2022"]}]


def test_palette_numbers_match_the_research_validator():
    # visual-quality.md §2.2.4: editorial deutan dE 3.8, clean protan dE 2.0.
    de = delta_e(simulate(parse_colour("#6B8E23"), "deutan"), simulate(parse_colour("#C0504D"), "deutan"))
    assert round(de, 1) == 3.8
    de = delta_e(simulate(parse_colour("#9AC05F"), "protan"), simulate(parse_colour("#D9B44A"), "protan"))
    assert round(de, 1) == 2.0
    assert round(contrast((255, 255, 255), (0, 0, 0)), 1) == 21.0
    assert categorical(["#1F5EFF"], "#FFFFFF") is None


def test_docx_page_estimate(tmp_path):
    docx = pytest.importorskip("docx")
    short, long_ = docx.Document(), docx.Document()
    for _ in range(5):
        short.add_paragraph("A paragraph of ordinary length that says one thing about the quarter.")
    for _ in range(80):
        long_.add_paragraph("A paragraph of ordinary length that says one thing about the quarter.")
    a, b = tmp_path / "a.docx", tmp_path / "b.docx"
    short.save(str(a))
    long_.save(str(b))
    assert measure.measure_docx(a)["pages_est"] == 1
    assert measure.measure_docx(b)["pages_est"] >= 2
